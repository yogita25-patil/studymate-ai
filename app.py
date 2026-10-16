import streamlit as st

# Optional browser camera + microphone support for Virtual Interview.
# Install with: python -m pip install streamlit-webrtc av
try:
    from streamlit_webrtc import webrtc_streamer, WebRtcMode, RTCConfiguration
    from aiortc.contrib.media import MediaRecorder
    from aiortc import MediaStreamTrack
    WEBRTC_AVAILABLE = True
except ImportError:
    WEBRTC_AVAILABLE = False
    webrtc_streamer = None
    WebRtcMode = None
    RTCConfiguration = None
    MediaRecorder = None
    MediaStreamTrack = None

import sys

import os
import io
import json
import inspect
import uuid
import time
import threading
from datetime import datetime

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
if PROJECT_DIR not in sys.path:
    sys.path.insert(0, PROJECT_DIR)


# ============================================================
# HR RESUME HELPER
# ============================================================

def _extract_resume_text(uploaded_resume):
    """Extract text from an uploaded HR resume without changing the existing RAG flow."""
    if uploaded_resume is None:
        return ""

    file_name = uploaded_resume.name.lower()
    file_bytes = uploaded_resume.getvalue()

    if file_name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()

    if file_name.endswith(".docx"):
        from docx import Document
        document = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in document.paragraphs).strip()

    return ""


# ============================================================
# VIRTUAL INTERVIEW MEDIA SETTINGS
# ============================================================

VIDEO_INTERVIEW_MEDIA_CONSTRAINTS = {
    "video": {
        "width": {"ideal": 1280, "max": 1920},
        "height": {"ideal": 720, "max": 1080},
        "frameRate": {"ideal": 30, "max": 30},
        "facingMode": "user",
    },
    "audio": {
        "echoCancellation": True,
        "noiseSuppression": True,
        "autoGainControl": True,
        "channelCount": 1,
    },
}


if WEBRTC_AVAILABLE and MediaStreamTrack is not None:

    class _GatedTrack(MediaStreamTrack):
        """
        Consume live camera/mic frames during the device check, but only
        pass frames to MediaRecorder after the final interview starts.

        This keeps the WebRTC connection alive during the pre-check and
        avoids creating a second recorder/peer connection at final start.
        """

        def __init__(self, source, gate):
            super().__init__()
            self.kind = source.kind
            self._source = source
            self._gate = gate

        async def recv(self):
            while True:
                frame = await self._source.recv()
                if self._gate.is_set():
                    return frame
                # Keep consuming frames so the browser/WebRTC connection
                # remains healthy while the user is checking devices.

        def stop(self):
            try:
                self._source.stop()
            except Exception:
                pass
            super().stop()


    class _GatedMediaRecorder(MediaRecorder):
        """MediaRecorder that records only after a shared Event is set."""

        def __init__(self, file_path, gate):
            self._gate = gate
            super().__init__(file_path)

        def addTrack(self, track):
            gated_track = _GatedTrack(track, self._gate)
            return super().addTrack(gated_track)

else:
    _GatedMediaRecorder = None


# ============================================================

# ADD SRC FOLDER TO PYTHON PATH

# ============================================================



SRC_PATH = os.path.join(

    os.path.dirname(__file__),

    "src"

)



if SRC_PATH not in sys.path:

    sys.path.append(SRC_PATH)





# ============================================================

# IMPORT PROJECT MODULES

# ============================================================

from flashcard_history import (
    save_flashcard_session,
    get_flashcard_history,
    get_flashcard_summary
)

from document_loader import extract_text_from_pdf

from text_splitter import split_text

from flashcard_generator import generate_flashcards

from vector_store import (
    add_documents,
    collection,
    get_all_documents,
    get_collection_count
)



from rag_qa import ask_question

from quiz_generator import generate_quiz
from source_ingestion import fetch_website, fetch_online_pdf
from interview_practice import generate_interview, evaluate_interview, save_interview_result, get_interview_history



from history_manager import (

    save_quiz_result,

    get_all_history,

    get_performance_summary

)



from image_manager import (
    save_image,
    explain_image,
    get_all_images,
    get_image
)

# ============================================================
# VECTOR STORE COMPATIBILITY HELPER
# ============================================================

def _call_add_documents(chunks, document_id, document_name, document_type):
    """Call vector_store.add_documents using only parameters supported by the installed version."""
    aliases = {
        "chunks": chunks,
        "document_id": document_id,
        "document_name": document_name,
        "document_type": document_type,
        "metadata": {"document_id": document_id, "document_name": document_name, "document_type": document_type},
    }
    sig = inspect.signature(add_documents)
    params = sig.parameters
    kwargs = {name: aliases[name] for name in params if name in aliases and params[name].kind in (inspect.Parameter.POSITIONAL_OR_KEYWORD, inspect.Parameter.KEYWORD_ONLY)}
    # If the function accepts **kwargs, provide the canonical metadata fields too.
    if any(p.kind == inspect.Parameter.VAR_KEYWORD for p in params.values()):
        kwargs.update({k: aliases[k] for k in ("document_id", "document_name", "document_type")})
    if "chunks" in params:
        return add_documents(**kwargs)
    # Fallback for positional-only/legacy implementations.
    return add_documents(chunks, document_id)


# ============================================================
# IMAGE MANAGER HELPERS
# ============================================================

def _extract_saved_image(record):
    """
    Extract image bytes, original filename and MIME type
    from the result returned by get_image().
    """

    if record is None:
        return None, None, None

    # If get_image() directly returns bytes
    if isinstance(record, (bytes, bytearray)):
        return bytes(record), None, None

    # If get_image() returns a dictionary
    if isinstance(record, dict):

        image_bytes = (
            record.get("image_bytes")
            or record.get("image_data")
            or record.get("data")
            or record.get("content")
            or record.get("file_bytes")
        )

        image_name = (
            record.get("original_name")
            or record.get("filename")
            or record.get("file_name")
            or record.get("name")
        )

        mime_type = (
            record.get("mime_type")
            or record.get("content_type")
            or record.get("media_type")
        )

        return image_bytes, image_name, mime_type

    return None, None, None




# ============================================================
# IMAGE MANAGER COMPATIBILITY HELPERS
# ============================================================

IMAGE_NAME_MAP_FILE = os.path.join(PROJECT_DIR, "image_filename_map.json")

def _load_image_name_map():
    try:
        if os.path.exists(IMAGE_NAME_MAP_FILE):
            with open(IMAGE_NAME_MAP_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                return data if isinstance(data, dict) else {}
    except Exception:
        pass
    return {}

def _save_image_name_map(data):
    try:
        with open(IMAGE_NAME_MAP_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
    except Exception:
        pass

def _call_save_image(image_bytes, original_name, mime_type):
    aliases = {
        "image_id": str(uuid.uuid4()),
        "image_bytes": image_bytes, "image_data": image_bytes, "data": image_bytes,
        "content": image_bytes, "file_bytes": image_bytes,
        "original_name": original_name, "filename": original_name, "file_name": original_name, "name": original_name,
        "mime_type": mime_type, "content_type": mime_type, "media_type": mime_type,
    }
    sig = inspect.signature(save_image)
    params = sig.parameters
    kwargs = {n: aliases[n] for n in params if n in aliases and params[n].kind in (inspect.Parameter.POSITIONAL_OR_KEYWORD, inspect.Parameter.KEYWORD_ONLY)}
    try:
        result = save_image(**kwargs) if kwargs else save_image(image_bytes)
    except TypeError:
        result = save_image(image_bytes)
    image_id = result.get("image_id") if isinstance(result, dict) else result
    if image_id:
        name_map = _load_image_name_map()
        name_map[str(image_id)] = {"original_name": original_name, "mime_type": mime_type or "image/png"}
        _save_image_name_map(name_map)
    return result


def _call_explain_image(image_bytes, mime_type, instruction):
    aliases = {
        "image_bytes": image_bytes, "image_data": image_bytes, "data": image_bytes,
        "content": image_bytes, "file_bytes": image_bytes,
        "mime_type": mime_type, "content_type": mime_type, "media_type": mime_type,
        "instruction": instruction, "prompt": instruction, "question": instruction,
    }
    sig = inspect.signature(explain_image)
    params = sig.parameters
    kwargs = {n: aliases[n] for n in params if n in aliases and params[n].kind in (inspect.Parameter.POSITIONAL_OR_KEYWORD, inspect.Parameter.KEYWORD_ONLY)}
    try:
        return explain_image(**kwargs) if kwargs else explain_image(image_bytes, instruction)
    except TypeError:
        return explain_image(image_bytes, instruction)

# ============================================================

# PAGE CONFIGURATION

# ============================================================



st.set_page_config(

    page_title="StudyMate AI",

    page_icon="🎓",

    layout="wide"

)

# ============================================================
# STUDYMATE AI - GLOBAL UI THEME
# UI ONLY: existing application logic/data flow is intentionally preserved.
# ============================================================

st.markdown(
    """
    <style>
    /* ---------- APP BACKGROUND ---------- */
    .stApp {
        background:
            radial-gradient(circle at 8% 4%, rgba(99,102,241,.10), transparent 28%),
            radial-gradient(circle at 92% 8%, rgba(14,165,233,.09), transparent 25%),
            linear-gradient(135deg, #f8fbff 0%, #f7f5ff 48%, #ffffff 100%);
    }

    /* ---------- SIDEBAR ---------- */
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #101a3a 0%, #172554 55%, #1e1b4b 100%);
        border-right: 0;
    }
    [data-testid="stSidebar"] * {
        color: #eef2ff;
    }
    [data-testid="stSidebar"] hr {
        border-color: rgba(255,255,255,.16);
    }
    [data-testid="stSidebar"] label,
    [data-testid="stSidebar"] p {
        color: #dbe4ff !important;
    }

    /* ---------- TYPOGRAPHY ---------- */
    .main-title {
        font-size: clamp(2.8rem, 5vw, 4.65rem);
        font-weight: 900;
        line-height: .98;
        letter-spacing: -2.8px;
        color: #ffffff;
        margin: .35rem 0 .8rem;
    }
    .hero-subtitle {
        font-size: 1.08rem;
        color: #dbe7ff;
        max-width: 760px;
        line-height: 1.7;
    }

    /* ---------- HOME HERO ---------- */
    .hero-box {
        position: relative;
        overflow: hidden;
        padding: 3.4rem 3.2rem;
        border-radius: 30px;
        background:
            radial-gradient(circle at 90% 20%, rgba(56,189,248,.28), transparent 28%),
            radial-gradient(circle at 15% 100%, rgba(129,140,248,.30), transparent 34%),
            linear-gradient(135deg, #111b41 0%, #312e81 52%, #4338ca 100%);
        border: 1px solid rgba(255,255,255,.16);
        box-shadow: 0 24px 55px rgba(30,41,90,.20);
        margin-bottom: 1.5rem;
    }
    .hero-box:after {
        content: "✦";
        position: absolute;
        right: 7%;
        top: 13%;
        font-size: 7rem;
        color: rgba(255,255,255,.08);
        transform: rotate(12deg);
    }
    .hero-badge {
        display: inline-block;
        padding: .42rem .9rem;
        border-radius: 999px;
        background: rgba(255,255,255,.13);
        color: #e0e7ff;
        border: 1px solid rgba(255,255,255,.20);
        font-size: .78rem;
        font-weight: 800;
        letter-spacing: .8px;
    }

    /* ---------- HOME SECTION HEADINGS ---------- */
    .section-label {
        color: #635bdb;
        font-size: .78rem;
        font-weight: 900;
        letter-spacing: 1.5px;
        text-transform: uppercase;
        margin-bottom: .2rem;
    }
    .home-section-title {
        color: #18234b;
        font-size: 2rem;
        font-weight: 850;
        letter-spacing: -1px;
        margin-bottom: .25rem;
    }
    .home-section-subtitle {
        color: #68748d;
        font-size: .98rem;
        margin-bottom: 1.1rem;
    }

    /* ---------- METRICS ---------- */
    .metric-card {
        background: rgba(255,255,255,.90);
        border: 1px solid #e3e8f4;
        border-radius: 20px;
        padding: 1.25rem 1.3rem;
        min-height: 112px;
        box-shadow: 0 10px 28px rgba(31,41,75,.07);
        backdrop-filter: blur(8px);
    }
    .metric-icon { font-size: 1.6rem; }
    .metric-label {
        color: #6b7690;
        font-size: .86rem;
        margin-top: .3rem;
    }
    .metric-value {
        color: #202d5b;
        font-size: 1.7rem;
        font-weight: 850;
    }

    /* ---------- FEATURE CARDS ---------- */
    .feature-card {
        background: rgba(255,255,255,.94);
        border: 1px solid #e4e8f2;
        border-radius: 21px;
        padding: 1.45rem;
        min-height: 180px;
        box-shadow: 0 10px 28px rgba(31,41,75,.055);
        transition: transform .18s ease, box-shadow .18s ease;
    }
    .feature-icon {
        display: inline-flex;
        width: 46px;
        height: 46px;
        align-items: center;
        justify-content: center;
        border-radius: 14px;
        background: linear-gradient(135deg, #eef2ff, #e0f2fe);
        font-size: 1.55rem;
        margin-bottom: .7rem;
    }
    .feature-title {
        color: #24315d;
        font-size: 1.05rem;
        font-weight: 800;
        margin-bottom: .35rem;
    }
    .feature-text {
        color: #69748a;
        line-height: 1.55;
        font-size: .9rem;
    }

    /* ---------- WORKFLOW ---------- */
    .step-card {
        background: linear-gradient(145deg, #ffffff, #f8faff);
        border: 1px solid #e5e9f3;
        border-radius: 19px;
        padding: 1.2rem;
        min-height: 150px;
        box-shadow: 0 8px 22px rgba(35,48,91,.05);
    }
    .step-number {
        color: #6257df;
        font-weight: 900;
        font-size: .76rem;
        letter-spacing: 1px;
        text-transform: uppercase;
    }
    .step-title {
        color: #28345c;
        font-weight: 800;
        margin: .45rem 0;
    }
    .step-text {
        color: #6b7488;
        font-size: .88rem;
        line-height: 1.5;
    }

    /* ---------- CONTROLS ---------- */
    div.stButton > button {
        border-radius: 11px;
        border: 1px solid #d8def0;
        font-weight: 700;
        min-height: 2.6rem;
    }
    div.stButton > button[kind="primary"] {
        background: linear-gradient(135deg,#5b5eea,#7c3aed);
        color: #fff;
        border: none;
        box-shadow: 0 7px 18px rgba(91,94,234,.18);
    }
    div[data-testid="stTextInput"] input,
    div[data-testid="stTextArea"] textarea,
    div[data-testid="stSelectbox"] div[data-baseweb="select"] > div {
        border-radius: 10px;
    }
    [data-testid="stFileUploader"] {
        border-radius: 16px;
    }
    [data-testid="stAlert"] {
        border-radius: 13px;
    }
    .footer-note {
        text-align: center;
        color: #8a93a6;
        font-size: .82rem;
        padding: 1.2rem 0 .5rem;
    }
    </style>
    """,
    unsafe_allow_html=True
)

# ============================================================
# REUSABLE PAGE INTRODUCTION
# ============================================================

def _page_intro(icon, title, kicker, description, tip):
    st.markdown(
        f"""
        <div style="
            position:relative;
            overflow:hidden;
            padding:1.65rem 1.8rem;
            margin-bottom:1.1rem;
            border-radius:24px;
            background:
                radial-gradient(circle at 92% 15%, rgba(56,189,248,.18), transparent 24%),
                linear-gradient(135deg,#111b41 0%,#312e81 60%,#4338ca 100%);
            box-shadow:0 16px 38px rgba(30,41,90,.14);
        ">
            <div style="color:#a5b4fc;font-size:.73rem;font-weight:900;
                        letter-spacing:1.5px;text-transform:uppercase;">
                {kicker}
            </div>
            <div style="color:white;font-size:2rem;font-weight:850;
                        letter-spacing:-.8px;margin:.25rem 0 .45rem;">
                {icon} {title}
            </div>
            <div style="color:#dbe7ff;line-height:1.6;max-width:900px;">
                {description}
            </div>
            <div style="
                display:inline-block;
                margin-top:.8rem;
                padding:.42rem .72rem;
                border-radius:10px;
                background:rgba(255,255,255,.10);
                border:1px solid rgba(255,255,255,.13);
                color:#e0e7ff;
                font-size:.78rem;
            ">
                ✦ {tip}
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

# ============================================================

# SESSION STATE

# ============================================================



defaults = {



    "page": "🏠 Home",



    "selected_document_id": None,



    "selected_document_name": None,



    "uploaded_file_name": None,



    "quiz": None,



    "quiz_topic": "",



    "quiz_difficulty": "",



    "quiz_type": "",



    "quiz_submitted": False,



    "user_answers": {},



    "quiz_score": 0,



    "quiz_total": 0,



    "quiz_saved": False,
    "interview_questions": None,
    "interview_answers": {},
    "interview_result": None,
    "interview_topic": "",
    "interview_type": "",
    "interview_difficulty": "",
    "interview_submitted": False,
    # Virtual/video interview state
    "video_interview_questions": None,
    "video_interview_answers": {},
    "video_interview_topic": "",
    "video_interview_type": "Mixed",
    "video_interview_difficulty": "Medium",
    "video_interview_current": 0,
    "video_interview_started": False,
    "video_interview_started_at": None,
    "video_interview_time_limit": 10,
    "video_interview_completed": False,
    "video_interview_result": None,
    "video_interview_recording_path": None,
    "video_interview_recording_stop": False,
    "video_interview_pending_evaluation": False,
    "video_interview_device_checked": False,
    "video_interview_instructions_ack": False,
    "video_interview_recording_gate": None

}



for key, value in defaults.items():



    if key not in st.session_state:

        st.session_state[key] = value





# ============================================================

# HELPER FUNCTIONS

# ============================================================



DOCUMENT_NAME_MAP_FILE = os.path.join(PROJECT_DIR, "document_filename_map.json")

def _load_document_name_map():
    try:
        if os.path.exists(DOCUMENT_NAME_MAP_FILE):
            with open(DOCUMENT_NAME_MAP_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                return data if isinstance(data, dict) else {}
    except Exception:
        pass
    return {}

def _save_document_name(document_id, document_name, document_type=""):
    if not document_id or not document_name:
        return
    try:
        name_map = _load_document_name_map()
        name_map[str(document_id)] = {
            "document_name": str(document_name),
            "document_type": str(document_type or "")
        }
        with open(DOCUMENT_NAME_MAP_FILE, "w", encoding="utf-8") as f:
            json.dump(name_map, f, indent=2, ensure_ascii=False)
    except Exception:
        pass

def _is_uuid_document_name(value):
    """Return True when a value looks like an internal UUID, not a filename."""
    if not value:
        return False
    try:
        uuid.UUID(str(value).strip())
        return True
    except (ValueError, AttributeError, TypeError):
        return False


def get_documents_from_collection():
    """Return only user-facing study documents, never raw internal UUID names."""
    try:
        data = collection.get(include=["metadatas"])
        metadatas = data.get("metadatas", [])
        name_map = _load_document_name_map()
        documents = {}

        for metadata in metadatas:
            if not metadata:
                continue

            document_id = metadata.get("document_id")
            raw_name = metadata.get("document_name")
            document_type = metadata.get("document_type", "") or ""

            if not document_id:
                continue

            # 1. Persistent map is the most reliable display-name source.
            mapped = name_map.get(str(document_id), {})
            mapped_name = mapped.get("document_name") if isinstance(mapped, dict) else None
            document_name = mapped_name or raw_name

            # 2. During the current upload session, the real uploaded filename
            # is available even if an old Chroma record has a bad name.
            if (
                str(document_id) == str(st.session_state.get("selected_document_id"))
                and st.session_state.get("uploaded_file_name")
            ):
                document_name = st.session_state.get("uploaded_file_name")

            # 3. IMPORTANT: old records created by the previous implementation
            # stored only the UUID as document_name and had no document_type.
            # There is no filename to recover from Chroma for those records.
            # Do not show internal UUIDs in the UI. They are legacy/orphaned
            # records, not the user's real filename.
            if _is_uuid_document_name(document_name):
                if mapped_name and not _is_uuid_document_name(mapped_name):
                    document_name = mapped_name
                else:
                    continue

            if not document_name:
                continue

            documents[str(document_id)] = {
                "document_id": document_id,
                "document_name": str(document_name),
                "document_type": str(document_type)
            }

        return list(documents.values())
    except Exception:
        return []


def normalize_question_type(question_type):



    if not question_type:

        return "MCQ"



    value = str(

        question_type

    ).strip().lower()



    if value in [

        "scenario",

        "scenario based",

        "scenario-based"

    ]:



        return "Scenario-Based"



    if value == "theory":



        return "Theory"



    return "MCQ"





def clear_quiz():



    st.session_state.quiz = None



    st.session_state.quiz_topic = ""



    st.session_state.quiz_difficulty = ""



    st.session_state.quiz_type = ""



    st.session_state.quiz_submitted = False



    st.session_state.user_answers = {}



    st.session_state.quiz_score = 0



    st.session_state.quiz_total = 0



    st.session_state.quiz_saved = False





# ============================================================

# SIDEBAR
# ============================================================

with st.sidebar:
    st.markdown(
        """
        <div style="padding:0.3rem 0 0.7rem 0;">
            <div style="font-size:2rem;">🎓</div>
            <div style="font-size:1.55rem;font-weight:850;color:#25345f;">StudyMate AI</div>
            <div style="color:#69748a;line-height:1.45;margin-top:0.2rem;">Your AI Learning Companion</div>
        </div>
        """,
        unsafe_allow_html=True
    )
    st.divider()
    st.markdown('<div class="section-label">Learn & Practice</div>', unsafe_allow_html=True)
    page = st.radio(
        "Choose a feature:",
        [
            "🏠 Home",
            "📚 My Notes",
            "💬 Ask My Notes",
            "📝 Quizzes",
            "🧠 Flashcards",
            "🎤 Interview Practice",
            "📊 Performance",
            "🖼️ Image Explanation"
        ],
        key="navigation",
        disabled=st.session_state.get("video_interview_started", False)
    )
    st.session_state.page = page
    st.divider()
    st.markdown('<div class="section-label">Current Study Material</div>', unsafe_allow_html=True)
    documents = get_documents_from_collection()
    if documents:
        document_names = [document["document_name"] for document in documents]
        current_name = st.session_state.selected_document_name
        default_index = document_names.index(current_name) if current_name in document_names else 0
        selected_name = st.selectbox("Select document:", document_names, index=default_index, key="document_selector")
        selected_document = next((document for document in documents if document["document_name"] == selected_name), None)
        if selected_document:
            # Keep the UUID only for internal RAG operations.
            # Always use the original document name for user-facing text.
            selected_document_id = selected_document.get("document_id")
            selected_document_name = selected_document.get("document_name") or "Untitled document"

            st.session_state.selected_document_id = selected_document_id
            st.session_state.selected_document_name = selected_document_name

            st.success(f"Using: {selected_document_name}")
    else:
        st.info("No study material uploaded yet.")
    st.divider()
    st.markdown(
        """
        <div style="padding:0.75rem 0.2rem;color:#69748a;font-size:0.84rem;line-height:1.55;">
            <b style="color:#3c4a72;">🤖 AI-powered learning</b><br>
            Study from your own material using RAG + Gemini.
        </div>
        """,
        unsafe_allow_html=True
    )


# ============================================================
# HOME
# ============================================================

if st.session_state.page == "🏠 Home":
    # ---------- HERO ----------
    st.markdown(
        """
        <div class="hero-box">
            <div class="hero-badge">🎓 YOUR PERSONAL AI LEARNING COMPANION</div>
            <div class="main-title">Learn Smarter.<br>Grow Faster.</div>
            <div class="hero-subtitle">
                StudyMate AI turns your course material into an interactive learning
                experience — helping you understand, practice, revise, prepare for
                interviews and track your progress in one place.
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

    try:
        _home_documents = get_documents_from_collection()
    except Exception:
        _home_documents = []
    try:
        _home_chunks = int(get_collection_count())
    except Exception:
        _home_chunks = 0

    # ---------- QUICK STATS ----------
    stat1, stat2, stat3, stat4 = st.columns(4)
    with stat1:
        st.markdown(
            f'<div class="metric-card"><div class="metric-icon">📚</div>'
            f'<div class="metric-value">{len(_home_documents)}</div>'
            f'<div class="metric-label">Study Materials</div></div>',
            unsafe_allow_html=True
        )
    with stat2:
        st.markdown(
            '<div class="metric-card"><div class="metric-icon">🧠</div>'
            '<div class="metric-value">RAG</div>'
            '<div class="metric-label">Personal Knowledge Base</div></div>',
            unsafe_allow_html=True
        )
    with stat3:
        st.markdown(
            '<div class="metric-card"><div class="metric-icon">🤖</div>'
            '<div class="metric-value">AI</div>'
            '<div class="metric-label">Personal Study Assistant</div></div>',
            unsafe_allow_html=True
        )
    with stat4:
        st.markdown(
            '<div class="metric-card"><div class="metric-icon">🎯</div>'
            '<div class="metric-value">24×7</div>'
            '<div class="metric-label">Learning Support</div></div>',
            unsafe_allow_html=True
        )

    st.markdown("<br>", unsafe_allow_html=True)

    # ---------- FEATURES ----------
    st.markdown('<div class="section-label">ONE PLATFORM • MANY WAYS TO LEARN</div>', unsafe_allow_html=True)
    st.markdown('<div class="home-section-title">Everything You Need to Learn Better</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="home-section-subtitle">Read less passively. Practice more actively. '
        'Use your own study material to make every session meaningful.</div>',
        unsafe_allow_html=True
    )

    feature_rows = [
        [
            ("📚", "My Notes", "Upload and organize your PDFs, text files and course material."),
            ("💬", "Ask My Notes", "Ask questions and get answers grounded in your selected study material."),
            ("📝", "Quizzes", "Generate MCQ, theory and scenario-based practice questions.")
        ],
        [
            ("🧠", "Flashcards", "Turn important concepts into quick revision cards and review them anytime."),
            ("🎤", "Interview Practice", "Practice technical interviews and receive an overall evaluation."),
            ("📊", "Performance", "Review assessment history, scores and learning progress.")
        ],
        [
            ("🖼️", "Image Explanation", "Understand diagrams, screenshots, handwritten notes and technical images."),
            ("🌐", "Online Learning", "Bring public website or online PDF content into your study workflow."),
            ("🤖", "AI-Powered", "Use RAG, embeddings and Gemini to make learning interactive.")
        ]
    ]

    for row in feature_rows:
        cols = st.columns(3, gap="medium")
        for col, (icon, title, description) in zip(cols, row):
            with col:
                st.markdown(
                    f'<div class="feature-card">'
                    f'<div class="feature-icon">{icon}</div>'
                    f'<div class="feature-title">{title}</div>'
                    f'<div class="feature-text">{description}</div>'
                    f'</div>',
                    unsafe_allow_html=True
                )
        st.markdown("<br>", unsafe_allow_html=True)

    # ---------- WORKFLOW ----------
    st.markdown('<div class="section-label">YOUR LEARNING JOURNEY</div>', unsafe_allow_html=True)
    st.markdown('<div class="home-section-title">From Material to Mastery</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="home-section-subtitle">A simple four-step flow designed around how students actually learn.</div>',
        unsafe_allow_html=True
    )

    steps = [
        ("01", "Upload", "Add your course material to your personal knowledge base."),
        ("02", "Understand", "Ask questions and retrieve relevant information from your material."),
        ("03", "Practice", "Generate quizzes, flashcards and interview questions."),
        ("04", "Improve", "Review results and focus on the topics that need more revision.")
    ]

    cols = st.columns(4, gap="medium")
    for col, (number, title, description) in zip(cols, steps):
        with col:
            st.markdown(
                f'<div class="step-card">'
                f'<div class="step-number">STEP {number}</div>'
                f'<div class="step-title">{title}</div>'
                f'<div class="step-text">{description}</div>'
                f'</div>',
                unsafe_allow_html=True
            )

    # ---------- BOTTOM VALUE BANNER ----------
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown(
        """
        <div class="hero-box" style="padding:1.8rem 2rem;">
            <div class="hero-badge">✨ STUDY DIFFERENTLY</div>
            <div style="font-size:1.7rem;font-weight:850;color:#fff;margin:.45rem 0 .4rem;">
                From passive reading to active learning.
            </div>
            <div style="color:#dbe7ff;line-height:1.65;max-width:900px;">
                StudyMate AI brings studying, question answering, assessment, revision,
                interview preparation and performance tracking into one student-focused workspace.
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="footer-note">StudyMate AI • Learn smarter. Practice better. Grow faster.</div>',
        unsafe_allow_html=True
    )


# ============================================================
# MY NOTES

# ============================================================



elif st.session_state.page == "📚 My Notes":



    st.title("📚 My Notes")



    st.write(

        """

        Upload your study material.

        Your previous documents will remain stored

        in your StudyMate AI library.

        """

    )



    st.divider()



    uploaded_file = st.file_uploader(

        "Upload your study material",

        type=[

            "pdf",

            "txt",

            "docx"

        ],

        key="study_material_uploader"

    )



    if uploaded_file is not None:



        st.success(

            f"Uploaded: {uploaded_file.name}"

        )



        if st.button(

            "📥 Process Study Material",

            type="primary"

        ):



            try:



                with st.spinner(

                    "Processing your study material..."

                ):



                    file_name = uploaded_file.name



                    extension = (

                        os.path.splitext(

                            file_name

                        )[1]

                        .lower()

                    )



                    # --------------------------------------------

                    # PDF

                    # --------------------------------------------



                    if extension == ".pdf":



                        text = extract_text_from_pdf(

                            uploaded_file

                        )



                        document_type = "PDF"



                    # --------------------------------------------

                    # TXT

                    # --------------------------------------------



                    elif extension == ".txt":



                        text = uploaded_file.read().decode(

                            "utf-8"

                        )



                        document_type = "TXT"



                    # --------------------------------------------

                    # DOCX

                    # --------------------------------------------



                    elif extension == ".docx":



                        from docx import Document



                        document = Document(

                            uploaded_file

                        )



                        text = "\n".join(

                            paragraph.text

                            for paragraph

                            in document.paragraphs

                        )



                        document_type = "DOCX"



                    else:



                        st.error(

                            "Unsupported file type."

                        )



                        text = ""



                        document_type = ""



                    # --------------------------------------------

                    # CHECK TEXT

                    # --------------------------------------------



                    if not text.strip():



                        st.error(

                            """

                            Could not extract text from

                            this file.

                            """

                        )



                    else:



                        # ----------------------------------------

                        # SPLIT TEXT

                        # ----------------------------------------



                        chunks = split_text(

                            text

                        )



                        # ----------------------------------------

                        # STORE DOCUMENT

                        # ----------------------------------------



                        document_id = str(uuid.uuid4())
                        _call_add_documents(
                            chunks,
                            document_id,
                            file_name,
                            document_type
                        )
                        _save_document_name(document_id, file_name, document_type)

                        # document_id is internal only; file_name is the display name.
                        display_document_name = file_name



                        # ----------------------------------------

                        # SAVE CURRENT DOCUMENT

                        # ----------------------------------------



                        st.session_state.selected_document_id = (

                            document_id

                        )



                        st.session_state.selected_document_name = (

                            display_document_name

                        )



                        st.session_state.uploaded_file_name = (

                            file_name

                        )



                        st.success(

                            "✅ Study material stored successfully!"

                        )



                        st.info(

                            f"""

                            **Document:** {file_name}



                            **Text extracted:** {len(text)} characters



                            **Chunks created:** {len(chunks)}



                            Your document is now available

                            in the StudyMate AI library.

                            """

                        )



                        st.rerun()



            except Exception as e:



                st.error(

                    f"Something went wrong: {e}"

                )



    st.divider()

    # ========================================================
    # ONLINE STUDY SOURCES
    # ========================================================

    st.subheader("🌐 Add Online Study Source")
    st.write("Add a public website or online PDF to your StudyMate knowledge base.")

    source_tab1, source_tab2 = st.tabs(["🌐 Website URL", "📄 Online PDF URL"])

    with source_tab1:
        website_url = st.text_input(
            "Website URL",
            placeholder="https://example.com/machine-learning-notes",
            key="website_source_url"
        )

        if st.button("🌐 Add Website", type="primary", key="add_website_source"):
            if not website_url.strip():
                st.warning("Please enter a website URL.")
            else:
                try:
                    with st.spinner("Reading website content..."):
                        title, source_text = fetch_website(website_url.strip())
                        # Keep online pages reasonably sized for faster embeddings and RAG.
                        source_text = source_text[:50000]
                        chunks = split_text(source_text)
                        document_name = f"Website — {title}"
                        document_id = str(uuid.uuid4())
                        _call_add_documents(
                            chunks,
                            document_id,
                            document_name,
                            "WEBSITE"
                        )
                        _save_document_name(document_id, document_name, "WEBSITE")
                        st.session_state.selected_document_id = document_id
                        st.session_state.selected_document_name = document_name

                    st.success(f"✅ Website added successfully: {title}")
                    st.info(f"{len(source_text):,} characters extracted • {len(chunks)} chunks created")
                    st.rerun()
                except Exception as e:
                    st.error(f"Could not add website: {e}")

    with source_tab2:
        pdf_url = st.text_input(
            "Online PDF URL",
            placeholder="https://example.com/notes.pdf",
            key="online_pdf_source_url"
        )

        if st.button("📄 Add Online PDF", type="primary", key="add_online_pdf_source"):
            if not pdf_url.strip():
                st.warning("Please enter an online PDF URL.")
            else:
                try:
                    with st.spinner("Downloading and reading PDF..."):
                        file_name, source_text = fetch_online_pdf(pdf_url.strip())
                        # Keep very large online PDFs manageable for faster processing.
                        source_text = source_text[:60000]
                        chunks = split_text(source_text)
                        document_name = f"Online PDF — {file_name}"
                        document_id = str(uuid.uuid4())
                        _call_add_documents(
                            chunks,
                            document_id,
                            document_name,
                            "ONLINE_PDF"
                        )
                        _save_document_name(document_id, document_name, "ONLINE_PDF")
                        st.session_state.selected_document_id = document_id
                        st.session_state.selected_document_name = document_name

                    st.success(f"✅ Online PDF added successfully: {file_name}")
                    st.info(f"{len(source_text):,} characters extracted • {len(chunks)} chunks created")
                    st.rerun()
                except Exception as e:
                    st.error(f"Could not add online PDF: {e}")

    st.divider()

    # --------------------------------------------------------
    # EXISTING DOCUMENTS
    # --------------------------------------------------------


    # --------------------------------------------------------



    st.subheader(

        "📚 Your Study Material Library"

    )



    documents = get_documents_from_collection()



    if documents:



        for document in documents:



            document_name = document.get(

                "document_name",

                "Unknown"

            )



            document_type = document.get(

                "document_type",

                ""

            )



            st.write(

                f"📄 **{document_name}** "

                f"({document_type})"

            )



    else:



        st.info(

            "No documents stored yet."

        )

# ============================================================

# ASK MY NOTES

# ============================================================



elif st.session_state.page == "💬 Ask My Notes":



    st.title("💬 Ask My Notes")



    st.write(

        """

        Ask questions based on the study material

        you have selected.

        """

    )



    # --------------------------------------------------------

    # CHECK DOCUMENT

    # --------------------------------------------------------



    selected_document_id = (

        st.session_state.selected_document_id

    )



    selected_document_name = (

        st.session_state.selected_document_name

    )

    # Never expose the internal UUID in the UI. Resolve the original name
    # from the current document library if an older session contains a UUID.
    if selected_document_id and (
        not selected_document_name
        or selected_document_name == selected_document_id
    ):
        for _doc in get_documents_from_collection():
            if _doc.get("document_id") == selected_document_id:
                selected_document_name = _doc.get("document_name") or "Untitled document"
                st.session_state.selected_document_name = selected_document_name
                break



    if not selected_document_id:



        st.warning(

            """

            ⚠️ Please upload a study document

            and select it from the sidebar first.

            """

        )



    else:



        st.info(

            f"📖 Studying from: **{selected_document_name}**"

        )



        question = st.text_input(

            "Enter your question:",

            placeholder=(

                "Example: What is a WHERE clause?"

            )

        )



        explanation_style = st.selectbox(

            "Answer style:",

            [

                "Simple Explanation",

                "Detailed Explanation",

                "Interview Explanation"

            ]

        )



        if st.button(

            "🔍 Ask StudyMate",

            type="primary"

        ):



            if not question.strip():



                st.warning(

                    "Please enter a question."

                )



            else:



                try:



                    with st.spinner(

                        "Searching your selected notes..."

                    ):



                        if (

                            explanation_style

                            == "Simple Explanation"

                        ):



                            style_instruction = """

Explain the answer in simple

student-friendly language.



Use short points and a simple example

when possible.

"""



                        elif (

                            explanation_style

                            == "Detailed Explanation"

                        ):



                            style_instruction = """

Give a detailed explanation using

the selected study material.



Explain step-by-step and include

an example when possible.

"""



                        else:



                            style_instruction = """

Explain the answer from an

interview perspective.



Use natural language that a student

can speak confidently in an interview.

"""



                        final_question = f"""

Question:

{question}



Answer Style:

{explanation_style}



Instructions:

{style_instruction}



IMPORTANT:

Answer ONLY using the selected

study material whenever possible.

"""



                        answer = ask_question(

                            final_question,

                            document_id=selected_document_id

                        )



                    st.subheader(

                        "🤖 StudyMate AI"

                    )



                    st.write(

                        answer

                    )



                except Exception as e:



                    st.error(

                        f"Something went wrong: {e}"

                    )





# ============================================================

# QUIZ GENERATOR

# ============================================================



elif st.session_state.page == "📝 Quizzes":



    st.title("📝 Quizzes")



    st.write(

        """

        Generate a personalized quiz from

        your selected study material.

        """

    )



    # --------------------------------------------------------

    # CHECK DOCUMENT

    # --------------------------------------------------------



    selected_document_id = (

        st.session_state.selected_document_id

    )



    selected_document_name = (

        st.session_state.selected_document_name

    )



    if not selected_document_id:



        st.warning(

            """

            ⚠️ Please upload and select a study

            document before generating a quiz.

            """

        )



    else:



        st.info(

            f"📖 Quiz based on: **{selected_document_name}**"

        )



        st.divider()



        # ----------------------------------------------------

        # QUIZ SETTINGS

        # ----------------------------------------------------



        col1, col2 = st.columns(2)



        with col1:



            topic = st.selectbox(

                "📚 Choose Topic",

                [

                    "Python",

                    "SQL",

                    "Machine Learning",

                    "Deep Learning",

                    "Power BI",

                    "Data Science",

                    "Excel",

                    "Statistics",

                    "GenAI"

                ]

            )



            difficulty = st.selectbox(

                "🎯 Difficulty",

                [

                    "Easy",

                    "Medium",

                    "Hard"

                ]

            )



        with col2:



            question_type = st.selectbox(

                "📝 Question Type",

                [

                    "MCQ",

                    "Theory",

                    "Scenario-Based"

                ]

            )



            number_option = st.selectbox(

                "🔢 Number of Questions",

                [

                    "5",

                    "10",

                    "15",

                    "20",

                    "Custom"

                ]

            )



            if number_option == "Custom":



                num_questions = st.number_input(

                    "Enter number of questions:",

                    min_value=1,

                    max_value=50,

                    value=5,

                    step=1

                )



            else:



                num_questions = int(

                    number_option

                )



        st.divider()



        # ----------------------------------------------------

        # QUESTION TYPE INFORMATION

        # ----------------------------------------------------



        if question_type == "MCQ":



            st.info(

                """

                **MCQ:** Choose the correct answer

                from the available options.

                """

            )



        elif question_type == "Theory":



            st.info(

                """

                **Theory:** Write the answer in

                your own words. After submission,

                StudyMate AI will show the

                reference answer.

                """

            )



        else:



            st.info(

                """

                **Scenario-Based:** You will receive

                a real-life situation. Explain what

                you would do, which concept you would

                use, and why.

                """

            )



        # ----------------------------------------------------

        # GENERATE QUIZ

        # ----------------------------------------------------



        if st.button(

            "🚀 Generate Quiz",

            type="primary"

        ):



            try:



                with st.spinner(

                    "Generating quiz from the selected document..."

                ):



                    quiz = generate_quiz(

                        topic=topic,

                        difficulty=difficulty,

                        num_questions=num_questions,

                        question_type=question_type,

                        document_id=selected_document_id

                    )



                if not quiz:



                    st.error(

                        """

                        No questions were generated.

                        Please try again.

                        """

                    )



                else:



                    st.session_state.quiz = quiz



                    st.session_state.quiz_topic = (

                        topic

                    )



                    st.session_state.quiz_difficulty = (

                        difficulty

                    )



                    st.session_state.quiz_type = (

                        question_type

                    )



                    st.session_state.quiz_submitted = (

                        False

                    )



                    st.session_state.user_answers = {}



                    st.session_state.quiz_score = 0



                    st.session_state.quiz_total = (

                        len(quiz)

                    )



                    st.session_state.quiz_saved = (

                        False

                    )



                    st.success(

                        f"✅ {len(quiz)} questions generated!"

                    )



            except Exception as e:



                st.error(

                    f"""

                    Something went wrong while

                    generating the quiz:



                    {e}

                    """

                )



        # ====================================================

        # DISPLAY QUIZ

        # ====================================================



        if st.session_state.quiz is not None:



            quiz = st.session_state.quiz



            st.divider()



            st.header(

                f"🎯 {st.session_state.quiz_topic} Quiz"

            )



            st.caption(

                f"Difficulty: "

                f"{st.session_state.quiz_difficulty}"

                f" | Type: "

                f"{st.session_state.quiz_type}"

            )



            st.divider()



            # ------------------------------------------------

            # DISPLAY EACH QUESTION

            # ------------------------------------------------



            for i, question in enumerate(quiz):



                # IMPORTANT:

                # This fixes the question_number error.



                question_number = i + 1



                question_text = question.get(

                    "question",

                    ""

                )



                current_type = normalize_question_type(

                    question.get(

                        "type",

                        st.session_state.quiz_type

                    )

                )



                # ------------------------------------------------

                # QUESTION

                # ------------------------------------------------



                st.subheader(

                    f"Q{question_number}. {question_text}"

                )



                st.caption(

                    f"Question Type: {current_type}"

                )



                # ------------------------------------------------

                # MCQ

                # ------------------------------------------------



                if current_type == "MCQ":



                    options = question.get(

                        "options",

                        []

                    )



                    if options:



                        selected_answer = st.radio(

                            "Choose your answer:",

                            options,

                            key=f"answer_{i}",

                            index=None

                        )



                        st.session_state.user_answers[i] = (

                            selected_answer

                        )



                    else:



                        st.warning(

                            "No options were generated for this question."

                        )



                        st.session_state.user_answers[i] = ""



                # ------------------------------------------------

                # THEORY

                # ------------------------------------------------



                elif current_type == "Theory":



                    theory_answer = st.text_area(

                        "✍️ Your answer:",

                        placeholder=(

                            "Explain the concept in your own words..."

                        ),

                        key=f"answer_{i}",

                        height=150

                    )



                    st.session_state.user_answers[i] = (

                        theory_answer

                    )



                # ------------------------------------------------

                # SCENARIO

                # ------------------------------------------------



                elif current_type == "Scenario-Based":



                    st.info(

                        """

                        💡 **How to answer:**



                        Explain:

                        1. What you would do.

                        2. Which concept or approach you would use.

                        3. Why you would choose it.

                        4. Give an example if possible.

                        """

                    )



                    scenario_answer = st.text_area(

                        "🎯 Your response:",

                        placeholder=(

                            "Explain what you would do, "

                            "which concept you would use, "

                            "and why..."

                        ),

                        key=f"answer_{i}",

                        height=180

                    )



                    st.session_state.user_answers[i] = (

                        scenario_answer

                    )



                st.divider()



            # ====================================================

            # SUBMIT QUIZ

            # ====================================================



            if not st.session_state.quiz_submitted:



                if st.button(

                    "✅ Submit Quiz",

                    type="primary"

                ):



                    score = 0



                    unanswered = 0



                    # ------------------------------------------------

                    # CHECK ANSWERS

                    # ------------------------------------------------



                    for i, question in enumerate(quiz):



                        current_type = normalize_question_type(

                            question.get(

                                "type",

                                st.session_state.quiz_type

                            )

                        )



                        user_answer = (

                            st.session_state.user_answers.get(

                                i,

                                ""

                            )

                        )



                        # --------------------------------------------

                        # MCQ SCORING

                        # --------------------------------------------



                        if current_type == "MCQ":



                            correct_answer = question.get(

                                "answer",

                                ""

                            )



                            if (

                                user_answer

                                and

                                str(user_answer).strip()

                                ==

                                str(correct_answer).strip()

                            ):



                                score += 1



                            elif not user_answer:



                                unanswered += 1



                        # --------------------------------------------

                        # THEORY / SCENARIO

                        # --------------------------------------------



                        else:



                            if (

                                not user_answer

                                or

                                not str(user_answer).strip()

                            ):



                                unanswered += 1



                    st.session_state.quiz_score = score



                    st.session_state.quiz_submitted = True



                    # ------------------------------------------------

                    # SAVE HISTORY

                    # ------------------------------------------------



                    if not st.session_state.quiz_saved:



                        try:



                            save_quiz_result(

                                topic=st.session_state.quiz_topic,

                                difficulty=st.session_state.quiz_difficulty,

                                num_questions=len(quiz),

                                score=score,

                                questions=quiz,

                                user_answers=st.session_state.user_answers

                            )



                            st.session_state.quiz_saved = True



                        except Exception as e:



                            st.warning(

                                f"""

                                Quiz completed, but history

                                could not be saved:



                                {e}

                                """

                            )



                    st.rerun()



            # ====================================================

            # RESULTS

            # ====================================================



            if st.session_state.quiz_submitted:



                st.divider()



                st.header(

                    "🏆 Quiz Results"

                )



                # ------------------------------------------------

                # COUNT MCQ QUESTIONS

                # ------------------------------------------------



                mcq_questions = 0



                for question in quiz:



                    q_type = normalize_question_type(

                        question.get(

                            "type",

                            st.session_state.quiz_type

                        )

                    )



                    if q_type == "MCQ":



                        mcq_questions += 1



                # ------------------------------------------------

                # SCORE

                # ------------------------------------------------



                if mcq_questions > 0:



                    percentage = round(

                        (

                            st.session_state.quiz_score

                            /

                            mcq_questions

                        )

                        * 100,

                        2

                    )



                    col1, col2 = st.columns(2)



                    with col1:



                        st.metric(

                            "MCQ Score",

                            f"{st.session_state.quiz_score}/{mcq_questions}"

                        )



                    with col2:



                        st.metric(

                            "Performance",

                            f"{percentage}%"

                        )



                    st.progress(

                        min(

                            percentage / 100,

                            1.0

                        )

                    )



                else:



                    st.info(

                        """

                        Your Theory/Scenario answers

                        have been recorded.



                        Compare your responses with

                        the reference answers below.

                        """

                    )



                # ------------------------------------------------

                # REVIEW

                # ------------------------------------------------



                st.subheader(

                    "📋 Review Your Answers"

                )



                for i, question in enumerate(quiz):



                    q_type = normalize_question_type(

                        question.get(

                            "type",

                            st.session_state.quiz_type

                        )

                    )



                    user_answer = (

                        st.session_state.user_answers.get(

                            i,

                            ""

                        )

                    )



                    correct_answer = question.get(

                        "answer",

                        ""

                    )



                    explanation = question.get(

                        "explanation",

                        ""

                    )



                    with st.expander(

                        f"Q{i + 1}. {q_type}"

                    ):



                        st.write(

                            f"**Question:** "

                            f"{question.get('question', '')}"

                        )



                        # ----------------------------------------

                        # MCQ REVIEW

                        # ----------------------------------------



                        if q_type == "MCQ":



                            st.write(

                                "**Your Answer:**"

                            )



                            if user_answer:



                                st.write(

                                    user_answer

                                )



                            else:



                                st.warning(

                                    "Not answered."

                                )



                            st.write(

                                "**Correct Answer:**"

                            )



                            st.success(

                                correct_answer

                            )



                            if (

                                user_answer

                                and

                                str(user_answer).strip()

                                ==

                                str(correct_answer).strip()

                            ):



                                st.success(

                                    "✅ Correct"

                                )



                            else:



                                st.error(

                                    "❌ Incorrect"

                                )



                        # ----------------------------------------

                        # THEORY / SCENARIO REVIEW

                        # ----------------------------------------



                        else:



                            st.write(

                                "**Your Answer:**"

                            )



                            if user_answer:



                                st.write(

                                    user_answer

                                )



                            else:



                                st.warning(

                                    "No answer provided."

                                )



                            st.write(

                                "**Reference Answer:**"

                            )



                            st.info(

                                correct_answer

                            )



                        # ----------------------------------------

                        # EXPLANATION

                        # ----------------------------------------



                        if explanation:



                            st.write(

                                "**Explanation:**"

                            )



                            st.write(

                                explanation

                            )



                st.divider()



                if st.button(

                    "🔄 Create New Quiz"

                ):



                    clear_quiz()



                    st.rerun()

    # ========================================================
    # QUIZ HISTORY
    # ========================================================

    st.divider()
    st.header("📜 Quiz History")

    try:
        history = get_all_history()
    except Exception as e:
        history = []
        st.warning(
            f"Could not load quiz history: {e}"
        )

    if not history:
        st.info(
            "You have not completed any quizzes yet."
        )
    else:
        st.write(
            f"Total quiz attempts: **{len(history)}**"
        )

        st.divider()

        for quiz_result in reversed(history):

            history_topic = quiz_result.get(
                "topic",
                "Unknown"
            )

            history_difficulty = quiz_result.get(
                "difficulty",
                "Unknown"
            )

            history_score = quiz_result.get(
                "score",
                0
            )

            history_total = quiz_result.get(
                "num_questions",
                0
            )

            history_percentage = quiz_result.get(
                "percentage",
                0
            )

            history_date = quiz_result.get(
                "date",
                ""
            )

            with st.expander(
                f"📚 {history_topic} | "
                f"{history_difficulty} | "
                f"{history_percentage}% | "
                f"{history_date}"
            ):

                col1, col2, col3 = st.columns(3)

                with col1:
                    st.metric(
                        "Score",
                        f"{history_score}/{history_total}"
                    )

                with col2:
                    st.metric(
                        "Percentage",
                        f"{history_percentage}%"
                    )

                with col3:
                    st.metric(
                        "Difficulty",
                        history_difficulty
                    )

                st.write(
                    f"**Topic:** {history_topic}"
                )

                st.write(
                    f"**Date:** {history_date}"
                )


# ============================================================
# FLASHCARDS
# ============================================================

elif st.session_state.page == "🧠 Flashcards":

    st.title("🧠 Flashcards")

    st.write(
        "Revise important concepts from your selected study material."
    )

    # ========================================================
    # SESSION STATE
    # ========================================================

    if "flashcards" not in st.session_state:
        st.session_state.flashcards = None

    if "flashcard_index" not in st.session_state:
        st.session_state.flashcard_index = 0

    if "flashcard_revealed" not in st.session_state:
        st.session_state.flashcard_revealed = False

    if "flashcard_status" not in st.session_state:
        st.session_state.flashcard_status = {}

    if "flashcard_session_completed" not in st.session_state:
        st.session_state.flashcard_session_completed = False

    if "flashcard_history_saved" not in st.session_state:
        st.session_state.flashcard_history_saved = False

    # ========================================================
    # USE THE CURRENT STUDY MATERIAL FROM THE SIDEBAR
    # ========================================================

    selected_document_id = st.session_state.get(
        "selected_document_id"
    )

    selected_document_name = st.session_state.get(
        "selected_document_name"
    )

    if selected_document_id and (
        not selected_document_name
        or selected_document_name == selected_document_id
    ):
        for _doc in get_documents_from_collection():
            if _doc.get("document_id") == selected_document_id:
                selected_document_name = _doc.get("document_name") or "Untitled document"
                st.session_state.selected_document_name = selected_document_name
                break

    if not selected_document_id:

        st.warning(
            "📚 Please upload and process a study document first."
        )

    else:

        st.info(
            f"📖 Flashcards based on: **{selected_document_name}**"
        )

        st.divider()

        # ====================================================
        # GENERATION SETTINGS
        # ====================================================

        topic = st.text_input(
            "📖 Topic",
            placeholder=(
                "Example: Machine Learning, SQL Joins, "
                "Deep Learning..."
            ),
            key="flashcard_topic"
        )

        col1, col2 = st.columns(2)

        with col1:

            difficulty = st.selectbox(
                "🎯 Difficulty",
                [
                    "Easy",
                    "Medium",
                    "Hard"
                ],
                key="flashcard_difficulty"
            )

        with col2:

            number_option = st.selectbox(
                "🔢 Number of Cards",
                [
                    "5",
                    "10",
                    "15",
                    "Custom"
                ],
                key="flashcard_number_option"
            )

            if number_option == "Custom":

                num_cards = st.number_input(
                    "Number of Cards",
                    min_value=1,
                    max_value=50,
                    value=10,
                    step=1,
                    key="flashcard_custom_number"
                )

            else:

                num_cards = int(number_option)

        if st.button(
            "🚀 Generate Flashcards",
            type="primary",
            key="generate_flashcards_button"
        ):

            if not topic.strip():

                st.warning(
                    "Please enter a topic first."
                )

            else:

                try:

                    with st.spinner(
                        "Creating your flashcards..."
                    ):

                        generated_cards = generate_flashcards(
                            topic=topic,
                            difficulty=difficulty,
                            num_cards=int(num_cards),
                            document_id=selected_document_id
                        )

                    st.session_state.flashcards = generated_cards
                    st.session_state.flashcard_index = 0
                    st.session_state.flashcard_revealed = False
                    st.session_state.flashcard_status = {}
                    st.session_state.flashcard_session_completed = False
                    st.session_state.flashcard_history_saved = False
                    st.session_state.flashcard_topic_value = topic
                    st.session_state.flashcard_difficulty_value = difficulty
                    st.session_state.flashcard_document_name_value = selected_document_name

                    st.success(
                        f"✅ {len(generated_cards)} flashcards generated!"
                    )

                except Exception as e:

                    st.error(
                        f"Something went wrong while generating flashcards: {e}"
                    )

        # ====================================================
        # ACTIVE FLASHCARD SESSION
        # ====================================================

        if (
            st.session_state.flashcards
            and not st.session_state.flashcard_session_completed
        ):

            flashcards = st.session_state.flashcards
            current_index = st.session_state.flashcard_index
            total_cards = len(flashcards)

            # Safety check in case a stale index remains in session state.
            if current_index >= total_cards:
                current_index = total_cards - 1
                st.session_state.flashcard_index = current_index

            current_card = flashcards[current_index]

            st.divider()

            st.subheader(
                f"🃏 Card {current_index + 1} of {total_cards}"
            )

            st.progress(
                (current_index + 1) / total_cards
            )

            st.markdown(
                f"""
                <div style="
                    padding: 30px;
                    border-radius: 15px;
                    border: 1px solid #cccccc;
                    margin-top: 20px;
                    margin-bottom: 20px;
                ">
                    <h3>❓ Question</h3>
                    <p style="font-size: 22px; font-weight: 600;">
                        {current_card.get("question", "")}
                    </p>
                </div>
                """,
                unsafe_allow_html=True
            )

            if not st.session_state.flashcard_revealed:

                if st.button(
                    "👀 Reveal Answer",
                    key=f"reveal_{current_index}"
                ):

                    st.session_state.flashcard_revealed = True
                    st.rerun()

            else:

                st.success(
                    f"✅ Answer\n\n{current_card.get('answer', '')}"
                )

                explanation = current_card.get(
                    "explanation",
                    ""
                )

                if explanation:

                    st.info(
                        f"💡 Explanation\n\n{explanation}"
                    )

                st.write("**How well did you know this?**")

                current_status = st.session_state.flashcard_status.get(
                    current_index
                )

                col1, col2 = st.columns(2)

                with col1:

                    if st.button(
                        "✅ I Know This",
                        key=f"known_{current_index}"
                    ):

                        st.session_state.flashcard_status[current_index] = "Known"

                        if current_index < total_cards - 1:
                            st.session_state.flashcard_index += 1
                            st.session_state.flashcard_revealed = False

                        else:
                            st.session_state.flashcard_session_completed = True

                        st.rerun()

                with col2:

                    if st.button(
                        "🔄 Need Revision",
                        key=f"revision_{current_index}"
                    ):

                        st.session_state.flashcard_status[current_index] = "Revision"

                        if current_index < total_cards - 1:
                            st.session_state.flashcard_index += 1
                            st.session_state.flashcard_revealed = False

                        else:
                            st.session_state.flashcard_session_completed = True

                        st.rerun()

                if current_status == "Known":
                    st.success("✅ Marked as Known")
                elif current_status == "Revision":
                    st.warning("🔄 Marked for Revision")

            # =================================================
            # NAVIGATION
            # =================================================

            st.divider()

            col1, col2, col3 = st.columns(3)

            with col1:

                if st.button(
                    "⬅️ Previous",
                    disabled=current_index == 0,
                    key="flashcard_previous"
                ):

                    st.session_state.flashcard_index -= 1
                    st.session_state.flashcard_revealed = False
                    st.rerun()

            with col2:

                if st.button(
                    "🔄 Restart",
                    key="flashcard_restart"
                ):

                    st.session_state.flashcard_index = 0
                    st.session_state.flashcard_revealed = False
                    st.session_state.flashcard_status = {}
                    st.session_state.flashcard_session_completed = False
                    st.session_state.flashcard_history_saved = False
                    st.rerun()

            with col3:

                if current_index < total_cards - 1:

                    if st.button(
                        "Next ➡️",
                        key="flashcard_next"
                    ):

                        st.session_state.flashcard_index += 1
                        st.session_state.flashcard_revealed = False
                        st.rerun()

                else:

                    all_marked = len(
                        st.session_state.flashcard_status
                    ) == total_cards

                    if st.button(
                        "🏁 Finish Session",
                        type="primary",
                        disabled=not all_marked,
                        key="finish_flashcard_session"
                    ):

                        st.session_state.flashcard_session_completed = True
                        st.rerun()

                    if not all_marked:

                        st.caption(
                            "Mark the last card as Known or Need Revision to finish."
                        )

        # ====================================================
        # SESSION RESULT
        # ====================================================

        if (
            st.session_state.flashcards
            and st.session_state.flashcard_session_completed
        ):

            flashcards = st.session_state.flashcards
            total_cards = len(flashcards)
            statuses = st.session_state.flashcard_status

            known = sum(
                1
                for status in statuses.values()
                if status == "Known"
            )

            revision = sum(
                1
                for status in statuses.values()
                if status == "Revision"
            )

            unanswered = total_cards - known - revision

            progress = round(
                (known / total_cards) * 100,
                2
            ) if total_cards else 0

            if not st.session_state.flashcard_history_saved:

                try:

                    save_flashcard_session(
                        topic=st.session_state.get(
                            "flashcard_topic_value",
                            topic
                        ),
                        difficulty=st.session_state.get(
                            "flashcard_difficulty_value",
                            difficulty
                        ),
                        document_name=st.session_state.get(
                            "flashcard_document_name_value",
                            selected_document_name
                        ),
                        total_cards=total_cards,
                        known=known,
                        revision=revision
                    )

                    st.session_state.flashcard_history_saved = True

                except Exception as e:

                    st.warning(
                        f"Session completed, but history could not be saved: {e}"
                    )

            st.divider()

            st.success(
                "🎉 Flashcard Session Completed!"
            )

            st.subheader(
                "📊 Your Session Result"
            )

            col1, col2, col3, col4 = st.columns(4)

            with col1:
                st.metric("Total Cards", total_cards)

            with col2:
                st.metric("✅ Known", known)

            with col3:
                st.metric("🔄 Revision", revision)

            with col4:
                st.metric("📈 Progress", f"{progress}%")

            if unanswered > 0:
                st.info(
                    f"ℹ️ {unanswered} card(s) were not marked."
                )

            st.progress(progress / 100)

            st.divider()

            col1, col2, col3 = st.columns(3)

            with col1:

                if st.button(
                    "🔄 Restart Session",
                    key="restart_completed_session"
                ):

                    st.session_state.flashcard_index = 0
                    st.session_state.flashcard_revealed = False
                    st.session_state.flashcard_status = {}
                    st.session_state.flashcard_session_completed = False
                    st.session_state.flashcard_history_saved = False
                    st.rerun()

            with col2:

                revision_indexes = [
                    index
                    for index, status in statuses.items()
                    if status == "Revision"
                ]

                if st.button(
                    "🔄 Review Revision Cards",
                    disabled=len(revision_indexes) == 0,
                    key="review_revision_cards"
                ):

                    st.session_state.flashcards = [
                        flashcards[index]
                        for index in revision_indexes
                    ]
                    st.session_state.flashcard_index = 0
                    st.session_state.flashcard_revealed = False
                    st.session_state.flashcard_status = {}
                    st.session_state.flashcard_session_completed = False
                    st.session_state.flashcard_history_saved = True
                    st.rerun()

            with col3:

                if st.button(
                    "🆕 Generate New Session",
                    key="new_flashcard_session"
                ):

                    st.session_state.flashcards = None
                    st.session_state.flashcard_index = 0
                    st.session_state.flashcard_revealed = False
                    st.session_state.flashcard_status = {}
                    st.session_state.flashcard_session_completed = False
                    st.session_state.flashcard_history_saved = False
                    st.rerun()

        # ====================================================
        # FLASHCARD HISTORY
        # ====================================================

        st.divider()
        st.header("📜 Flashcard History")

        try:
            history = get_flashcard_history()
            summary = get_flashcard_summary()
        except Exception as e:
            history = []
            summary = {}
            st.warning(
                f"Could not load flashcard history: {e}"
            )

        if history:

            st.subheader("📊 Overall Flashcard Statistics")

            col1, col2, col3, col4, col5 = st.columns(5)

            with col1:
                st.metric(
                    "Sessions",
                    summary.get("total_sessions", 0)
                )

            with col2:
                st.metric(
                    "Total Cards",
                    summary.get("total_cards", 0)
                )

            with col3:
                st.metric(
                    "✅ Known",
                    summary.get("known", 0)
                )

            with col4:
                st.metric(
                    "🔄 Revision",
                    summary.get("revision", 0)
                )

            with col5:
                st.metric(
                    "Average Progress",
                    f"{summary.get('average_progress', 0)}%"
                )

            st.divider()
            st.subheader("📚 Previous Sessions")

            for session in reversed(history):

                with st.expander(
                    f"📖 {session.get('topic', 'Unknown')} | "
                    f"{session.get('progress', 0)}% | "
                    f"{session.get('date', '')}"
                ):

                    col1, col2, col3, col4 = st.columns(4)

                    with col1:
                        st.metric(
                            "Cards",
                            session.get("total_cards", 0)
                        )

                    with col2:
                        st.metric(
                            "Known",
                            session.get("known", 0)
                        )

                    with col3:
                        st.metric(
                            "Revision",
                            session.get("revision", 0)
                        )

                    with col4:
                        st.metric(
                            "Progress",
                            f"{session.get('progress', 0)}%"
                        )

                    st.write(
                        f"**Document:** "
                        f"{session.get('document_name', 'Unknown')}"
                    )

                    st.write(
                        f"**Difficulty:** "
                        f"{session.get('difficulty', 'Unknown')}"
                    )

                    st.write(
                        f"**Date:** "
                        f"{session.get('date', '')}"
                    )

        else:

            st.info(
                "No flashcard sessions completed yet."
            )

# ============================================================
# INTERVIEW PRACTICE
# ============================================================

elif st.session_state.page == "🎤 Interview Practice":

    _page_intro(
        "🎤",
        "Interview Practice",
        "PREPARE • SPEAK • IMPROVE",
        "Practice technical, HR, mixed and project-based interviews using your own study material, then receive an overall evaluation.",
        "Complete the interview first → get the evaluation together."
    )

    st.divider()

    # Navigation is intentionally locked only after the final interview starts.
    interview_mode = st.radio(
        "Choose interview mode",
        ["📝 Text Interview", "🎥 Virtual Interview"],
        horizontal=True,
        key="interview_mode",
        disabled=st.session_state.get("video_interview_started", False)
    )

    # ========================================================
    # VIRTUAL INTERVIEW
    # ========================================================

    if interview_mode == "🎥 Virtual Interview":

        # --------------------------------------------------------
        # FINISH RECORDING + EVALUATE
        # --------------------------------------------------------
        if (
            st.session_state.get("video_interview_pending_evaluation", False)
            and st.session_state.get("video_interview_result") is None
        ):
            questions = st.session_state.get("video_interview_questions") or []
            answers = st.session_state.get("video_interview_answers") or {}

            # The WebRTC component disappeared on the previous rerun,
            # which tells streamlit-webrtc to stop the recorder. Give PyAV
            # a short time to finalize the WebM container.
            recording_path = st.session_state.get("video_interview_recording_path")
            if recording_path:
                for _ in range(20):
                    if os.path.exists(recording_path) and os.path.getsize(recording_path) > 0:
                        break
                    time.sleep(0.15)

            answered_indices = [
                i for i in range(len(questions))
                if str(answers.get(i, "")).strip()
            ]

            if not answered_indices:
                st.warning("Please answer at least one question before submitting.")
                st.session_state.video_interview_pending_evaluation = False
                st.session_state.video_interview_recording_stop = False
            else:
                answered_questions = [questions[i] for i in answered_indices]
                answered_answers = {
                    position: answers.get(original_index, "")
                    for position, original_index in enumerate(answered_indices)
                }

                try:
                    with st.spinner(
                        "StudyMate AI is evaluating your completed virtual interview..."
                    ):
                        evaluation = evaluate_interview(
                            questions=answered_questions,
                            answers=answered_answers,
                            topic=st.session_state.video_interview_topic,
                            interview_type=st.session_state.video_interview_type,
                            difficulty=st.session_state.video_interview_difficulty
                        )

                        save_interview_result(
                            topic=st.session_state.video_interview_topic,
                            interview_type=f"Virtual - {st.session_state.video_interview_type}",
                            difficulty=st.session_state.video_interview_difficulty,
                            document_name=st.session_state.get(
                                "selected_document_name", "Unknown"
                            ),
                            questions=answered_questions,
                            answers=answered_answers,
                            evaluation=evaluation
                        )

                    st.session_state.video_interview_result = evaluation
                    st.session_state.video_interview_completed = True
                    st.session_state.video_interview_pending_evaluation = False
                    st.session_state.video_interview_recording_stop = False
                    st.rerun()

                except Exception as e:
                    st.error(f"Could not evaluate virtual interview: {e}")
                    st.session_state.video_interview_pending_evaluation = False
                    st.session_state.video_interview_recording_stop = False

        # --------------------------------------------------------
        # COMPLETED INTERVIEW RESULT
        # --------------------------------------------------------
        elif st.session_state.get("video_interview_result") is not None:

            result = st.session_state.video_interview_result
            questions = st.session_state.get("video_interview_questions") or []
            answers = st.session_state.get("video_interview_answers") or {}

            st.success("🎉 Virtual Interview Completed!")
            st.subheader("📊 Interview Result")

            score = float(result.get("overall_score", 0))
            answered_count = sum(
                1 for value in answers.values() if str(value).strip()
            )

            col1, col2, col3 = st.columns(3)
            with col1:
                st.metric("Overall Score", f"{score:.1f}/10")
            with col2:
                st.metric("Answered", f"{answered_count}/{len(questions)}")
            with col3:
                st.metric(
                    "Topic",
                    st.session_state.get("video_interview_topic", "Interview")
                )

            st.progress(max(0.0, min(1.0, score / 10.0)))

            if result.get("summary"):
                st.info(result["summary"])

            col1, col2 = st.columns(2)
            with col1:
                st.subheader("💪 Strengths")
                strengths = result.get("strengths", [])
                if strengths:
                    for item in strengths:
                        st.write(f"• {item}")
                else:
                    st.write("No strengths were provided.")

            with col2:
                st.subheader("📈 Areas to Improve")
                improvements = result.get("improvements", [])
                if improvements:
                    for item in improvements:
                        st.write(f"• {item}")
                else:
                    st.write("No improvement points were provided.")

            recording_path = st.session_state.get("video_interview_recording_path")
            if recording_path and os.path.exists(recording_path):
                st.divider()
                st.subheader("🎥 Interview Recording")
                st.caption("Your final interview recording is saved locally for this StudyMate AI session.")
                try:
                    with open(recording_path, "rb") as video_file:
                        video_bytes = video_file.read()
                    st.video(video_bytes)
                    st.download_button(
                        "⬇️ Download Interview Recording",
                        data=video_bytes,
                        file_name=os.path.basename(recording_path),
                        mime="video/webm",
                        key="download_virtual_interview_recording"
                    )
                except Exception as e:
                    st.warning(f"Recording was created but could not be opened: {e}")

            st.divider()
            st.subheader("📝 Question-wise Feedback")
            feedback_items = result.get("questions", [])

            for i, question in enumerate(questions):
                with st.expander(f"Q{i + 1}. {question.get('question', '')}"):
                    st.write("**Your answer:**")
                    st.write(answers.get(i, "") or "No answer provided.")

                    if i < len(feedback_items):
                        item = feedback_items[i]
                        st.write(f"**Score:** {item.get('score', 0)}/10")
                        st.write(
                            f"**Feedback:** {item.get('feedback', 'No feedback provided.')}"
                        )
                        if item.get("better_answer"):
                            st.write("**Better answer:**")
                            st.write(item["better_answer"])

            if st.button(
                "🔄 Start New Virtual Interview",
                type="primary",
                key="new_virtual_interview"
            ):
                for key, value in {
                    "video_interview_questions": None,
                    "video_interview_answers": {},
                    "video_interview_topic": "",
                    "video_interview_type": "Mixed",
                    "video_interview_difficulty": "Medium",
                    "video_interview_current": 0,
                    "video_interview_started": False,
                    "video_interview_started_at": None,
                    "video_interview_completed": False,
                    "video_interview_result": None,
                    "video_interview_recording_path": None,
                    "video_interview_recording_stop": False,
                    "video_interview_pending_evaluation": False,
                    "video_interview_device_checked": False,
                    "video_interview_instructions_ack": False,
                    "video_interview_recording_gate": None,
                }.items():
                    st.session_state[key] = value
                st.rerun()

        else:
            documents = get_documents_from_collection()

            if not documents:
                st.warning("Please upload or add study material in My Notes first.")

            # ====================================================
            # SETUP SCREEN
            # ====================================================
            elif st.session_state.get("video_interview_questions") is None:
                st.subheader("🎥 Virtual Interview Setup")
                st.caption(
                    "Choose your material and configure the interview. Your camera and microphone will be checked before the final session starts."
                )

                video_type = st.selectbox(
                    "🎯 1. Choose Interview Type",
                    ["Technical", "HR", "Mixed", "Project Based"],
                    index=2,
                    key="video_interview_type_selector"
                )

                st.caption(
                    {
                        "Technical": "Technical concepts and scenario-based questions.",
                        "HR": "Role-focused HR preparation using your notes.",
                        "Mixed": "Technical + general interview practice.",
                        "Project Based": "Project explanation and implementation questions."
                    }[video_type]
                )

                document_names = [doc["document_name"] for doc in documents]
                current_name = st.session_state.get(
                    "selected_document_name", document_names[0]
                )
                if current_name not in document_names:
                    current_name = document_names[0]

                selected_name = st.selectbox(
                    "📚 2. Choose Study Material",
                    document_names,
                    index=document_names.index(current_name),
                    key="video_interview_document_selector"
                )

                selected_doc = next(
                    (doc for doc in documents if doc["document_name"] == selected_name),
                    None
                )
                if selected_doc:
                    st.session_state.selected_document_id = selected_doc["document_id"]
                    st.session_state.selected_document_name = selected_doc["document_name"]

                col1, col2 = st.columns(2)
                with col1:
                    video_topic = st.text_input(
                        "🎯 3. Interview Topic / Role",
                        value=st.session_state.get("video_interview_topic", ""),
                        placeholder=(
                            "Example: Data Scientist"
                            if video_type == "HR"
                            else "Example: Machine Learning"
                        )
                    )

                    video_hr_notes = ""
                    video_hr_resume = None
                    if video_type == "HR":
                        video_hr_notes = st.text_area(
                            "📝 HR Interview Notes",
                            placeholder="Add role details, job description or preparation notes.",
                            height=100,
                            key="video_hr_interview_notes"
                        )
                        video_hr_resume = st.file_uploader(
                            "📄 Upload Resume (optional)",
                            type=["pdf", "docx"],
                            help="Use your resume to tailor HR questions to your profile.",
                            key="video_hr_resume_uploader"
                        )
                        if video_hr_resume is not None:
                            st.success(f"Resume selected: {video_hr_resume.name}")

                with col2:
                    video_difficulty = st.selectbox(
                        "Difficulty",
                        ["Easy", "Medium", "Hard"],
                        index=1,
                        key="video_interview_difficulty_selector"
                    )
                    video_num_questions = st.number_input(
                        "Number of Questions",
                        min_value=3,
                        max_value=10,
                        value=5,
                        step=1,
                        key="video_interview_num_questions"
                    )
                    video_time_limit = st.selectbox(
                        "⏱️ Interview Time Limit",
                        [5, 10, 15, 20, 30],
                        index=1,
                        format_func=lambda x: f"{x} minutes",
                        key="video_interview_time_limit_selector"
                    )

                if st.button(
                    "🚀 Prepare Virtual Interview",
                    type="primary",
                    key="prepare_virtual_interview"
                ):
                    if not video_topic.strip():
                        st.warning("Please enter an interview topic or role.")
                    elif video_type == "HR" and not video_hr_notes.strip():
                        st.warning(
                            "For an HR interview, please add your interview notes, role details or job description."
                        )
                    elif not st.session_state.get("selected_document_id"):
                        st.warning("Please select study material first.")
                    else:
                        try:
                            video_generation_topic = video_topic.strip()

                            if video_type == "HR":
                                video_generation_topic = (
                                    f"{video_topic.strip()}\n\n"
                                    f"HR Interview Notes:\n{video_hr_notes.strip()}"
                                )
                                if video_hr_resume is not None:
                                    try:
                                        resume_text = _extract_resume_text(video_hr_resume)
                                        if resume_text:
                                            st.session_state.video_hr_resume_name = video_hr_resume.name
                                            st.session_state.video_hr_resume_text = resume_text
                                            video_generation_topic += (
                                                "\n\nResume Content:\n"
                                                + resume_text[:12000]
                                            )
                                        else:
                                            st.warning("Could not extract text from the uploaded resume. You can continue using the HR notes.")
                                    except Exception as resume_error:
                                        st.warning(f"Could not read the resume: {resume_error}")

                            with st.spinner(
                                "StudyMate AI is preparing your interview questions..."
                            ):
                                questions = generate_interview(
                                    topic=video_generation_topic,
                                    interview_type=video_type,
                                    difficulty=video_difficulty,
                                    num_questions=int(video_num_questions),
                                    document_id=st.session_state.selected_document_id
                                )

                            recording_dir = os.path.join(
                                os.path.dirname(__file__),
                                "interview_recordings"
                            )
                            os.makedirs(recording_dir, exist_ok=True)
                            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                            recording_path = os.path.join(
                                recording_dir,
                                f"virtual_interview_{stamp}.webm"
                            )

                            # The gate is shared by the recorder created during
                            # the camera check. It consumes camera frames during
                            # the check but only writes frames after final start.
                            st.session_state.video_interview_recording_gate = threading.Event()
                            st.session_state.video_interview_questions = questions
                            st.session_state.video_interview_answers = {}
                            st.session_state.video_interview_topic = video_topic.strip()
                            st.session_state.video_interview_type = video_type
                            st.session_state.video_interview_difficulty = video_difficulty
                            st.session_state.video_interview_current = 0
                            st.session_state.video_interview_started = False
                            st.session_state.video_interview_started_at = None
                            st.session_state.video_interview_completed = False
                            st.session_state.video_interview_time_limit = int(video_time_limit)
                            st.session_state.video_interview_result = None
                            st.session_state.video_interview_recording_path = recording_path
                            st.session_state.video_interview_recording_stop = False
                            st.session_state.video_interview_pending_evaluation = False
                            st.session_state.video_interview_device_checked = False
                            st.session_state.video_interview_instructions_ack = False
                            st.rerun()
                        except Exception as e:
                            st.error(f"Could not prepare virtual interview: {e}")

            # ====================================================
            # PRE-CHECK + INSTRUCTIONS + FINAL START
            # ====================================================
            elif not st.session_state.get("video_interview_started", False):
                questions = st.session_state.video_interview_questions

                st.subheader("🎥 Virtual Interview — Device Check")
                st.write(
                    "First check your camera and microphone. This is only a live preview; "
                    "the timer does not start until you click Start Final Interview Session."
                )

                if not WEBRTC_AVAILABLE or MediaRecorder is None:
                    st.error(
                        "Camera/microphone support is not installed. Activate your StudyMate environment and run:"
                    )
                    st.code("python -m pip install -U streamlit-webrtc av", language="bash")
                else:
                    recording_path = st.session_state.get("video_interview_recording_path")
                    recording_gate = st.session_state.get("video_interview_recording_gate")

                    if recording_gate is None:
                        recording_gate = threading.Event()
                        st.session_state.video_interview_recording_gate = recording_gate

                    rtc_config = RTCConfiguration(
                        {
                            "iceServers": [
                                {"urls": ["stun:stun.l.google.com:19302"]}
                            ]
                        }
                    )

                    def _gated_recorder_factory():
                        return _GatedMediaRecorder(
                            recording_path,
                            recording_gate
                        )

                    try:
                        ctx = webrtc_streamer(
                            key="studymate_virtual_interview_recorder",
                            mode=WebRtcMode.SENDRECV,
                            rtc_configuration=rtc_config,
                            media_stream_constraints=VIDEO_INTERVIEW_MEDIA_CONSTRAINTS,
                            in_recorder_factory=_gated_recorder_factory,
                            desired_playing_state=True,
                            media_stream_constraints_timeout=15,
                            video_html_attrs={
                                "controls": False,
                                "autoPlay": True,
                                "playsInline": True,
                                "muted": True,
                                "style": {
                                    "width": "100%",
                                    "maxWidth": "1100px",
                                    "margin": "0 auto",
                                    "borderRadius": "12px",
                                },
                            },
                            media_toggle_controls=False,
                        )
                    except TypeError:
                        # Compatibility with older streamlit-webrtc versions.
                        ctx = webrtc_streamer(
                            key="studymate_virtual_interview_recorder",
                            mode=WebRtcMode.SENDRECV,
                            rtc_configuration=rtc_config,
                            media_stream_constraints=VIDEO_INTERVIEW_MEDIA_CONSTRAINTS,
                            in_recorder_factory=_gated_recorder_factory,
                            desired_playing_state=True,
                        )

                    playing = bool(getattr(getattr(ctx, "state", None), "playing", False))
                    if playing:
                        st.success("🟢 Camera and microphone connected successfully.")
                        st.session_state.video_interview_device_checked = True
                    else:
                        st.info(
                            "Click the START button inside the camera preview and allow browser camera/microphone permissions."
                        )

                st.divider()
                st.subheader("📋 Interview Instructions")
                st.markdown(
                    """
                    - Keep your face clearly visible in the camera frame.
                    - Keep your microphone enabled and stay in a quiet environment.
                    - Answer naturally, as you would in a real interview.
                    - The timer starts only after **Start Final Interview Session**.
                    - Your camera and microphone remain connected throughout the session.
                    - You can move between questions and submit the interview at any point.
                    - The AI evaluates the completed answers together after submission.
                    """
                )

                st.checkbox(
                    "I have checked my camera, microphone and interview instructions.",
                    key="video_interview_instructions_ack",
                    disabled=not st.session_state.get("video_interview_device_checked", False)
                )

                col1, col2, col3 = st.columns(3)
                with col1:
                    st.metric("Questions", len(questions))
                with col2:
                    st.metric("Time", f"{st.session_state.video_interview_time_limit} min")
                with col3:
                    st.metric("Difficulty", st.session_state.video_interview_difficulty)

                can_start = (
                    WEBRTC_AVAILABLE
                    and st.session_state.get("video_interview_device_checked", False)
                    and st.session_state.get("video_interview_instructions_ack", False)
                )

                if st.button(
                    "🎬 Start Final Interview Session",
                    type="primary",
                    key="start_final_virtual_interview",
                    disabled=not can_start
                ):
                    gate = st.session_state.get("video_interview_recording_gate")
                    if gate is not None:
                        gate.set()

                    st.session_state.video_interview_started = True
                    st.session_state.video_interview_started_at = time.time()
                    st.session_state.video_interview_current = 0
                    st.session_state.video_interview_recording_stop = False
                    st.rerun()

                if not can_start:
                    st.caption(
                        "Complete the camera/microphone check and tick the instruction confirmation to enable the final start button."
                    )

            # ====================================================
            # ACTIVE VIRTUAL INTERVIEW
            # ====================================================
            else:
                questions = st.session_state.video_interview_questions or []
                current = st.session_state.video_interview_current
                current = max(0, min(current, len(questions) - 1))
                st.session_state.video_interview_current = current

                started_at = st.session_state.get("video_interview_started_at")
                total_seconds = int(st.session_state.video_interview_time_limit * 60)
                if started_at is None:
                    started_at = time.time()
                    st.session_state.video_interview_started_at = started_at

                remaining = max(0, total_seconds - int(time.time() - started_at))

                # ----------------------------------------------------
                # INTERVIEW HEADER
                # ----------------------------------------------------
                st.markdown(
                    f"## 🎥 Virtual Interview · Question {current + 1} of {len(questions)}"
                )
                st.info(questions[current].get("question", ""))

                timer_placeholder = st.empty()

                def _render_live_timer():
                    live_started = st.session_state.get("video_interview_started_at")
                    live_limit = int(
                        st.session_state.get("video_interview_time_limit", 10) * 60
                    )
                    if live_started is None:
                        return
                    live_remaining = max(
                        0,
                        live_limit - int(time.time() - live_started)
                    )
                    mm = live_remaining // 60
                    ss = live_remaining % 60
                    message = f"⏱️ Time remaining: {mm:02d}:{ss:02d}"
                    if live_remaining <= 60:
                        timer_placeholder.error(message)
                    else:
                        timer_placeholder.success(message)

                if hasattr(st, "fragment"):
                    @st.fragment(run_every="1s")
                    def _live_timer_fragment():
                        _render_live_timer()
                    _live_timer_fragment()
                else:
                    _render_live_timer()

                if remaining <= 0:
                    st.error("⏰ Time is up. Submit the answers you have completed.")

                # ----------------------------------------------------
                # CAMERA + MICROPHONE — SAME WEBRTC KEY AS PRE-CHECK
                # ----------------------------------------------------
                st.markdown("### 📹 Your Live Interview")
                st.caption(
                    "HD camera + microphone are connected continuously. The recorder started writing only when the final interview began."
                )

                if not WEBRTC_AVAILABLE or MediaRecorder is None:
                    st.error("Camera/microphone recording is not available.")
                else:
                    recording_path = st.session_state.get("video_interview_recording_path")
                    recording_gate = st.session_state.get("video_interview_recording_gate")

                    if recording_gate is None:
                        recording_gate = threading.Event()
                        recording_gate.set()
                        st.session_state.video_interview_recording_gate = recording_gate

                    rtc_config = RTCConfiguration(
                        {
                            "iceServers": [
                                {"urls": ["stun:stun.l.google.com:19302"]}
                            ]
                        }
                    )

                    def _gated_recorder_factory_active():
                        return _GatedMediaRecorder(
                            recording_path,
                            recording_gate
                        )

                    try:
                        webrtc_streamer(
                            key="studymate_virtual_interview_recorder",
                            mode=WebRtcMode.SENDRECV,
                            rtc_configuration=rtc_config,
                            media_stream_constraints=VIDEO_INTERVIEW_MEDIA_CONSTRAINTS,
                            in_recorder_factory=_gated_recorder_factory_active,
                            desired_playing_state=True,
                            media_stream_constraints_timeout=15,
                            video_html_attrs={
                                "controls": False,
                                "autoPlay": True,
                                "playsInline": True,
                                "muted": True,
                                "style": {
                                    "width": "100%",
                                    "maxWidth": "1100px",
                                    "margin": "0 auto",
                                    "borderRadius": "12px",
                                },
                            },
                            media_toggle_controls=False,
                        )
                    except TypeError:
                        webrtc_streamer(
                            key="studymate_virtual_interview_recorder",
                            mode=WebRtcMode.SENDRECV,
                            rtc_configuration=rtc_config,
                            media_stream_constraints=VIDEO_INTERVIEW_MEDIA_CONSTRAINTS,
                            in_recorder_factory=_gated_recorder_factory_active,
                            desired_playing_state=True,
                        )

                # ----------------------------------------------------
                # ANSWER TRANSCRIPT
                # ----------------------------------------------------
                st.markdown("### 🗣️ Your Answer")
                answer = st.text_area(
                    "Answer transcript",
                    value=st.session_state.video_interview_answers.get(current, ""),
                    key=f"video_interview_answer_{current}",
                    height=150,
                    placeholder=(
                        "Speak your answer naturally. Type or paste the transcript here; "
                        "the transcript is used for AI evaluation."
                    )
                )

                st.divider()

                # ----------------------------------------------------
                # CONTROLS
                # ----------------------------------------------------
                nav1, nav2, nav3, nav4 = st.columns(4)

                with nav1:
                    if st.button(
                        "⬅️ Previous",
                        disabled=current == 0,
                        key="video_previous"
                    ):
                        st.session_state.video_interview_answers[current] = answer
                        st.session_state.video_interview_current -= 1
                        st.rerun()

                with nav2:
                    if st.button(
                        "💾 Save Answer",
                        key="video_save_answer"
                    ):
                        st.session_state.video_interview_answers[current] = answer
                        st.success("Answer saved for this question.")

                with nav3:
                    if current < len(questions) - 1:
                        if st.button(
                            "Next ➡️",
                            type="primary",
                            key="video_next"
                        ):
                            st.session_state.video_interview_answers[current] = answer
                            st.session_state.video_interview_current += 1
                            st.rerun()
                    else:
                        st.button(
                            "Next ➡️",
                            disabled=True,
                            key="video_next_disabled"
                        )

                with nav4:
                    if st.button(
                        "🏁 Submit Interview",
                        type="primary",
                        key="submit_virtual_interview"
                    ):
                        st.session_state.video_interview_answers[current] = answer
                        st.session_state.video_interview_recording_stop = True
                        st.session_state.video_interview_pending_evaluation = True
                        st.session_state.video_interview_started = False
                        st.session_state.video_interview_started_at = None
                        st.rerun()

                if remaining <= 0:
                    st.warning("The time limit has been reached. Please submit now.")

                st.caption(
                    "You may submit after any question. Only answered questions are evaluated. "
                    "The final interview recording is saved locally as a WebM file."
                )

    # TEXT INTERVIEW — EXISTING FEATURE
    # ========================================================

    # ========================================================

    else:

        # ------------------------------------------------------------
        # Completed evaluation view
        # ------------------------------------------------------------
        if st.session_state.get("interview_result") is not None:

            result = st.session_state.interview_result

            st.success("🎉 Interview Completed!")
            st.subheader("📊 Your Interview Result")

            score = float(result.get("overall_score", 0))

            col1, col2, col3 = st.columns(3)

            with col1:
                st.metric("Overall Score", f"{score:.1f}/10")

            with col2:
                st.metric(
                    "Questions",
                    len(st.session_state.get("interview_questions") or [])
                )

            with col3:
                st.metric(
                    "Topic",
                    st.session_state.get("interview_topic", "Interview")
                )

            st.progress(max(0.0, min(1.0, score / 10.0)))

            if result.get("summary"):
                st.info(result["summary"])

            col1, col2 = st.columns(2)

            with col1:
                st.subheader("💪 Strengths")
                strengths = result.get("strengths", [])
                if strengths:
                    for item in strengths:
                        st.write(f"• {item}")
                else:
                    st.write("No strengths were provided.")

            with col2:
                st.subheader("📈 Areas to Improve")
                improvements = result.get("improvements", [])
                if improvements:
                    for item in improvements:
                        st.write(f"• {item}")
                else:
                    st.write("No improvement points were provided.")

            st.divider()
            st.subheader("📝 Question-wise Feedback")

            questions = st.session_state.get("interview_questions") or []
            answers = st.session_state.get("interview_answers") or {}
            feedback_items = result.get("questions", [])

            for i, question in enumerate(questions):
                with st.expander(
                    f"Q{i + 1}. {question.get('question', '')}"
                ):
                    st.write("**Your answer:**")
                    st.write(answers.get(i, "") or "No answer provided.")

                    if i < len(feedback_items):
                        item = feedback_items[i]
                        st.write(
                            f"**Score:** {item.get('score', 0)}/10"
                        )
                        st.write(
                            f"**Feedback:** {item.get('feedback', 'No feedback provided.')}"
                        )

                        if item.get("better_answer"):
                            st.write("**Better answer:**")
                            st.write(item["better_answer"])

            st.divider()

            col1, col2 = st.columns(2)

            with col1:
                if st.button(
                    "🔄 Start New Interview",
                    type="primary",
                    key="start_new_text_interview"
                ):
                    st.session_state.interview_questions = None
                    st.session_state.interview_answers = {}
                    st.session_state.interview_result = None
                    st.session_state.interview_topic = ""
                    st.session_state.interview_type = ""
                    st.session_state.interview_difficulty = ""
                    st.session_state.interview_submitted = False
                    st.rerun()

            with col2:
                if st.button(
                    "📜 View Interview History",
                    key="view_text_interview_history"
                ):
                    history = get_interview_history()

                    if history:
                        st.subheader("📜 Interview History")
                        for item in reversed(history):
                            st.write(
                                f"**{item.get('topic', 'Interview')}** | "
                                f"{item.get('interview_type', 'Mixed')} | "
                                f"{item.get('difficulty', 'Medium')} | "
                                f"Score: {item.get('score', 0)}/10 | "
                                f"{item.get('date', '')}"
                            )
                    else:
                        st.info("No interview history yet.")

        # ------------------------------------------------------------
        # Interview setup / answering
        # ------------------------------------------------------------
        else:

            documents = get_documents_from_collection()

            if not documents:
                st.warning(
                    "Please upload or add study material in My Notes first."
                )
            else:

                # Interview type comes first so the setup adapts to the student's goal.
                interview_type = st.selectbox(
                    "🎯 1. Choose Interview Type",
                    ["Technical", "HR", "Mixed", "Project Based"],
                    index=2,
                    key="interview_type_selector"
                )

                st.caption(
                    {
                        "Technical": "Concepts, definitions, why/how and technical scenarios.",
                        "HR": "HR-style questions based on your notes, role or preparation material.",
                        "Mixed": "A balanced combination of technical and general interview questions.",
                        "Project Based": "Questions focused on projects, implementation and your technical decisions."
                    }[interview_type]
                )

                document_names = [
                    doc["document_name"] for doc in documents
                ]

                current_name = st.session_state.get(
                    "selected_document_name",
                    document_names[0]
                )

                if current_name not in document_names:
                    current_name = document_names[0]

                selected_name = st.selectbox(
                    "📚 2. Choose Study Material",
                    document_names,
                    index=document_names.index(current_name),
                    key="interview_document_selector"
                )

                selected_doc = next(
                    (
                        doc for doc in documents
                        if doc["document_name"] == selected_name
                    ),
                    None
                )

                if selected_doc:
                    st.session_state.selected_document_id = selected_doc["document_id"]
                    st.session_state.selected_document_name = selected_doc["document_name"]

                col1, col2 = st.columns(2)

                with col1:
                    topic = st.text_input(
                        "🎯 3. Interview Topic / Role",
                        value=st.session_state.get("interview_topic", ""),
                        placeholder=(
                            "Example: Data Scientist"
                            if interview_type == "HR"
                            else "Example: SQL Joins"
                        ),
                        key="text_interview_topic"
                    )

                    interview_notes = ""
                    hr_resume = None
                    if interview_type == "HR":
                        interview_notes = st.text_area(
                            "📝 HR Interview Notes",
                            placeholder=(
                                "Add role details, job description, "
                                "company expectations or topics you want to practice."
                            ),
                            height=110,
                            key="hr_interview_notes"
                        )
                        hr_resume = st.file_uploader(
                            "📄 Upload Resume (optional)",
                            type=["pdf", "docx"],
                            help="Upload your resume so HR questions can be tailored to your education, skills, projects and experience.",
                            key="hr_resume_uploader"
                        )
                        if hr_resume is not None:
                            st.success(f"Resume selected: {hr_resume.name}")

                with col2:
                    difficulty = st.selectbox(
                        "Difficulty",
                        ["Easy", "Medium", "Hard"],
                        index=1,
                        key="interview_difficulty_selector"
                    )

                    num_questions = st.number_input(
                        "Number of Questions",
                        min_value=3,
                        max_value=15,
                        value=5,
                        step=1,
                        key="interview_num_questions"
                    )

                if st.session_state.get("interview_questions") is None:

                    if st.button(
                        "🎤 Start Interview",
                        type="primary",
                        key="start_text_interview"
                    ):

                        if not topic.strip():
                            st.warning(
                                "Please enter an interview topic or role."
                            )
                        elif interview_type == "HR" and not interview_notes.strip():
                            st.warning(
                                "For an HR interview, please add your interview notes, role details or job description."
                            )
                        else:
                            try:
                                generation_topic = topic.strip()

                                if interview_type == "HR":
                                    generation_topic = (
                                        f"{topic.strip()}\n\n"
                                        f"HR Interview Notes:\n{interview_notes.strip()}"
                                    )
                                    if hr_resume is not None:
                                        try:
                                            resume_text = _extract_resume_text(hr_resume)
                                            if resume_text:
                                                st.session_state.hr_resume_name = hr_resume.name
                                                st.session_state.hr_resume_text = resume_text
                                                generation_topic += (
                                                    "\n\nResume Content:\n"
                                                    + resume_text[:12000]
                                                )
                                            else:
                                                st.warning("Could not extract text from the uploaded resume. You can continue using the HR notes.")
                                        except Exception as resume_error:
                                            st.warning(f"Could not read the resume: {resume_error}")

                                with st.spinner(
                                    "StudyMate AI is preparing your interview..."
                                ):
                                    questions = generate_interview(
                                        topic=generation_topic,
                                        interview_type=interview_type,
                                        difficulty=difficulty,
                                        num_questions=int(num_questions),
                                        document_id=st.session_state.selected_document_id
                                    )

                                st.session_state.interview_questions = questions
                                st.session_state.interview_answers = {}
                                st.session_state.interview_result = None
                                st.session_state.interview_topic = topic.strip()
                                st.session_state.interview_type = interview_type
                                st.session_state.interview_difficulty = difficulty
                                st.session_state.interview_submitted = False
                                st.rerun()

                            except Exception as e:
                                st.error(
                                    f"Could not generate interview: {e}"
                                )

                else:

                    questions = st.session_state.interview_questions

                    st.info(
                        "Answer all questions first. Your answers will be evaluated together after you submit the complete interview."
                    )

                    for i, question in enumerate(questions):
                        st.subheader(
                            f"Q{i + 1}. {question.get('question', '')}"
                        )

                        st.session_state.interview_answers[i] = st.text_area(
                            "Your Answer",
                            value=st.session_state.interview_answers.get(i, ""),
                            key=f"interview_answer_{i}",
                            height=130,
                            placeholder="Type your answer here..."
                        )

                    st.divider()

                    if st.button(
                        "📊 Submit Complete Interview",
                        type="primary",
                        key="submit_text_interview"
                    ):

                        if any(
                            not str(
                                st.session_state.interview_answers.get(i, "")
                            ).strip()
                            for i in range(len(questions))
                        ):
                            st.warning(
                                "Please answer all interview questions before submitting."
                            )
                        else:

                            try:
                                with st.spinner(
                                    "StudyMate AI is evaluating your complete interview..."
                                ):
                                    evaluation = evaluate_interview(
                                        questions=questions,
                                        answers=st.session_state.interview_answers,
                                        topic=st.session_state.interview_topic,
                                        interview_type=st.session_state.interview_type,
                                        difficulty=st.session_state.interview_difficulty
                                    )

                                    save_interview_result(
                                        topic=st.session_state.interview_topic,
                                        interview_type=st.session_state.interview_type,
                                        difficulty=st.session_state.interview_difficulty,
                                        document_name=st.session_state.get(
                                            "selected_document_name",
                                            "Unknown"
                                        ),
                                        questions=questions,
                                        answers=st.session_state.interview_answers,
                                        evaluation=evaluation
                                    )

                                st.session_state.interview_result = evaluation
                                st.session_state.interview_submitted = True
                                st.rerun()

                            except Exception as e:
                                st.error(
                                    f"Could not evaluate interview: {e}"
                                )

# ============================================================

# PERFORMANCE

# ============================================================



elif st.session_state.page == "📊 Performance":



    st.title(

        "📊 Student Performance"

    )



    try:



        summary = get_performance_summary()



    except Exception as e:



        st.error(

            f"Could not load performance data: {e}"

        )



        summary = {}



    total_attempts = summary.get(

        "total_attempts",

        0

    )



    average_percentage = summary.get(

        "average_percentage",

        0

    )



    best_percentage = summary.get(

        "best_percentage",

        0

    )



    if total_attempts == 0:



        st.info(

            """

            No performance data yet.



            Complete a quiz to start

            tracking your performance.

            """

        )



    else:



        # --------------------------------------------------------

        # OVERALL PERFORMANCE

        # --------------------------------------------------------



        col1, col2, col3 = st.columns(3)



        with col1:



            st.metric(

                "📝 Total Attempts",

                total_attempts

            )



        with col2:



            st.metric(

                "📊 Average Score",

                f"{average_percentage}%"

            )



        with col3:



            st.metric(

                "🏆 Best Score",

                f"{best_percentage}%"

            )



        st.divider()



        # --------------------------------------------------------

        # SUBJECT PERFORMANCE

        # --------------------------------------------------------



        st.subheader(

            "📚 Subject-Wise Performance"

        )



        topics = summary.get(

            "topics",

            {}

        )



        if not topics:



            st.info(

                "No subject-wise performance available."

            )



        else:



            for topic, data in topics.items():



                attempts = data.get(

                    "attempts",

                    0

                )



                average = data.get(

                    "average",

                    0

                )



                best = data.get(

                    "best",

                    0

                )



                latest = data.get(

                    "latest",

                    0

                )



                st.markdown(

                    f"### 📖 {topic}"

                )



                col1, col2, col3, col4 = (

                    st.columns(4)

                )



                with col1:



                    st.metric(

                        "Attempts",

                        attempts

                    )



                with col2:



                    st.metric(

                        "Average",

                        f"{average}%"

                    )



                with col3:



                    st.metric(

                        "Best",

                        f"{best}%"

                    )



                with col4:



                    st.metric(

                        "Latest",

                        f"{latest}%"

                    )



                st.progress(

                    min(

                        average / 100,

                        1.0

                    )

                )



                if average >= 80:



                    st.success(

                        "🟢 Strong performance in this topic."

                    )



                elif average >= 60:



                    st.warning(

                        "🟡 This topic needs some revision."

                    )



                else:



                    st.error(

                        "🔴 Consider revising this topic."

                    )



                st.divider()

            # ============================================================

# IMAGE EXPLANATION

# ============================================================



elif st.session_state.page == "🖼️ Image Explanation":



    st.title("🖼️ Image Explanation")



    st.write(

        """

        Upload a blackboard photo, handwritten note, diagram,

        screenshot or any educational image and let StudyMate AI

        explain it for you.

        """

    )



    st.divider()



    # ========================================================

    # ========================================================
    # UPLOAD NEW IMAGE ONLY
    # ========================================================

    st.subheader("📤 Upload Image")
    st.caption("Upload a PNG, JPG, JPEG or WEBP educational image. The image can be explained immediately.")

    uploaded_image = st.file_uploader(
        "Choose your study image",
        type=["png", "jpg", "jpeg", "webp"],
        key="image_explanation_uploader"
    )

    selected_image_bytes = None
    selected_image_name = None
    selected_mime_type = None

    if uploaded_image is not None:
        selected_image_bytes = uploaded_image.getvalue()
        selected_image_name = uploaded_image.name
        selected_mime_type = uploaded_image.type or "image/png"
        st.image(selected_image_bytes, caption=f"Uploaded: {selected_image_name}", use_container_width=True)
        st.success(f"✅ {selected_image_name} uploaded successfully and is ready for explanation.")
        if st.button("💾 Save Image to StudyMate Library", type="secondary", key="save_uploaded_image"):
            try:
                _call_save_image(selected_image_bytes, selected_image_name, selected_mime_type)
                st.success("✅ Image saved successfully to your StudyMate image library.")
            except Exception as e:
                st.warning(f"Image explanation is still available. Library save skipped: {e}")

    # EXPLANATION OPTIONS

    # ========================================================



    if selected_image_bytes:



        st.divider()



        st.subheader(

            "🧠 What do you want StudyMate AI to do?"

        )



        explanation_type = st.selectbox(

            "Choose explanation style",

            [

                "Explain in Simple Points",

                "Explain in Detailed Points",

                "Explain Important Concepts",

                "Explain for Exam Preparation",

                "Explain for Interview Preparation"

            ]

        )



        # ----------------------------------------------------

        # CUSTOM QUESTION

        # ----------------------------------------------------



        custom_question = st.text_area(

            "💬 Ask something about this image (optional)",

            placeholder=(

                "Example: Explain this blackboard note "

                "in very simple language."

            )

        )



        # ----------------------------------------------------

        # BUILD INSTRUCTION

        # ----------------------------------------------------



        if custom_question.strip():



            instruction = custom_question



        elif explanation_type == "Explain in Simple Points":



            instruction = """

Explain this note in very simple student-friendly points.

Assume that the student is seeing this concept for the first time.

Give a small example wherever useful.

"""



        elif explanation_type == "Explain in Detailed Points":



            instruction = """

Explain the complete educational content in detail.

Break the explanation into clear sections and points.

Explain definitions, concepts, formulas, examples and diagrams.

"""



        elif explanation_type == "Explain Important Concepts":



            instruction = """

Identify the most important concepts from this image.

Explain each important concept clearly.

Also mention what the student should remember for exams.

"""



        elif explanation_type == "Explain for Exam Preparation":



            instruction = """

Explain this material from an exam preparation perspective.

Identify important definitions, concepts, formulas and possible

exam questions.

"""



        else:



            instruction = """

Explain this material from a technical interview perspective.

Identify important concepts and then provide possible interview

questions based on the image.

"""



        # ====================================================

        # EXPLAIN BUTTON

        # ====================================================



        if st.button(

            "🤖 Explain Selected Image with StudyMate AI",

            type="primary"

        ):



            try:



                with st.spinner(

                    "StudyMate AI is analyzing the image..."

                ):

                    answer = _call_explain_image(
                        selected_image_bytes,
                        selected_mime_type,
                        instruction
                    )



                st.divider()



                st.subheader(

                    "🤖 StudyMate AI Explanation"

                )



                st.markdown(

                    answer

                )



            except Exception as e:



                st.error(

                    f"Something went wrong while analyzing the image: {e}"

                )